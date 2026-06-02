"""
جميع مسارات الـ API — All API Routes
Blueprint واحد يجمع كل المسارات
"""
import datetime, io, json, os, smtplib, mimetypes, urllib.parse
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from flask import Blueprint, request, send_file, g, current_app, send_from_directory
from flask_jwt_extended import (
    create_access_token, create_refresh_token,
    jwt_required, get_jwt_identity,
)
from sqlalchemy import func
from werkzeug.security import check_password_hash
from werkzeug.utils import secure_filename
from app.models import (
    db, User, Warehouse, Category,     Supplier, Item,
    Stock, StockMovement, Transfer, InventoryCount,
    InventoryCountLine, Notification, Project,
    ItemAttachment, NotificationConfig, BackupConfig,
    ProjectInvoice,
    PurchaseRequest, PRItem, ApprovalLog,
    RFQ, RFQSupplier, Quotation, QuotationItem,
    PurchaseOrder, POItem,
    GoodsReceipt, GRNItem,
    PurchaseReturn, PReturnItem, SaleOrder, SaleItem,
    StockMovement,
    SupplierEvaluation, SupplierProfile, InventoryLayer,
    Lot, UnitConversion, ApprovalChain, ApprovalStep, ScheduledCount,
)

def send_notif_email(subject, body, to=None):
    try:
        host = NotificationConfig.query.filter_by(key="smtp_host").first()
        port = NotificationConfig.query.filter_by(key="smtp_port").first()
        user = NotificationConfig.query.filter_by(key="smtp_user").first()
        pwd  = NotificationConfig.query.filter_by(key="smtp_pass").first()
        frm  = NotificationConfig.query.filter_by(key="email_from").first()
        enb  = NotificationConfig.query.filter_by(key="email_enabled").first()
        if not host or not host.value or (enb and enb.value != "true"): return
        recipients = [to] if to else [u.email for u in User.query.filter(User.role.in_(["admin","super_admin"])).all() if u.email]
        for r in recipients:
            try:
                msg = MIMEMultipart()
                msg["From"] = frm.value if frm else user.value if user else "noreply@wms.local"
                msg["To"] = r; msg["Subject"] = subject
                msg.attach(MIMEText(body, "plain", "utf-8"))
                with smtplib.SMTP(host.value, int(port.value if port else 587), timeout=10) as s:
                    s.starttls()
                    if user and user.value and pwd and pwd.value: s.login(user.value, pwd.value)
                    s.send_message(msg)
            except: pass
    except: pass

def _uname(uid): u = User.query.get(uid); return u.name if u else "—"

from app.services import (
    AuthService, DashboardService, ItemService,
    MovementService, TransferService, CountService,
    SupplierService, UserService, ReportService,
    AuditService,
    StockInquiryService,
)
from app.utils import ok, created, err, not_found, forbidden, unauthorized, require_role, validate, gen_ref, parse_date, today_str
from app.limiter import limiter

api = Blueprint("api", __name__, url_prefix="/api")

MAX_FILE_SIZE = 16 * 1024 * 1024  # 16 MB
ALLOWED_EXTENSIONS = {"pdf","jpg","jpeg","png","gif","doc","docx","xls","xlsx","csv","txt","zip"}

# ══════════════════════════════════════════════════════════════
#  AUTH
# ══════════════════════════════════════════════════════════════
@api.route("/auth/login", methods=["POST"])
@limiter.limit("10 per minute")
def login():
    """
    User login
    ---
    parameters:
      - name: body
        in: body
        required: true
        schema:
          properties:
            username: {type: string, example: admin}
            password: {type: string, example: admin123}
    responses:
      200: {description: Login successful}
      401: {description: Invalid credentials}
    """
    data     = request.get_json() or {}
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""
    if not username or not password:
        return err("اسم المستخدم وكلمة المرور مطلوبان")
    user, error = AuthService.login(username, password)
    if error:
        return err(error, 401)
    access  = create_access_token(identity=str(user.id))
    refresh = create_refresh_token(identity=str(user.id))
    return ok({"access_token":access,"refresh_token":refresh,
               "user":user.to_dict(),"permissions":user.get_permissions()},
              "تم تسجيل الدخول")

@api.route("/auth/refresh", methods=["POST"])
@jwt_required(refresh=True)
@limiter.limit("20 per minute")
def refresh():
    uid  = int(get_jwt_identity())
    user = User.query.get(uid)
    if not user or not user.is_active:
        return unauthorized("الحساب غير صالح")
    return ok({"access_token": create_access_token(identity=str(uid))})

@api.route("/auth/logout", methods=["POST"])
@jwt_required()
def logout():
    uid  = int(get_jwt_identity())
    user = User.query.get(uid)
    if user:
        AuditService.log("logout","auth",uid,f"خروج: {user.name}")
        db.session.commit()
    return ok(message="تم تسجيل الخروج")

@api.route("/auth/me", methods=["GET"])
@jwt_required()
def me():
    user = User.query.get(get_jwt_identity())
    if not user:
        return not_found("المستخدم غير موجود")
    return ok({"user":user.to_dict(),"permissions":user.get_permissions()})

@api.route("/auth/change-password", methods=["POST"])
@jwt_required()
def change_password():
    uid = int(get_jwt_identity())
    data = request.get_json() or {}
    ok_flag, error = AuthService.change_password(
        uid, data.get("old_password",""), data.get("new_password","")
    )
    if error: return err(error)
    return ok(message="تم تغيير كلمة المرور")

# ══════════════════════════════════════════════════════════════
#  DASHBOARD
# ══════════════════════════════════════════════════════════════
@api.route("/dashboard", methods=["GET"])
@jwt_required()
def dashboard():
    """
    Dashboard KPIs
    ---
    responses:
      200: {description: Dashboard data with KPIs, warehouses, critical items}
    """
    uid = int(get_jwt_identity())
    data = DashboardService.get_data(uid)
    return ok(data)

# ══════════════════════════════════════════════════════════════
#  WAREHOUSES
# ══════════════════════════════════════════════════════════════
@api.route("/warehouses", methods=["GET"])
@jwt_required()
def get_warehouses():
    return ok([w.to_dict() for w in Warehouse.query.filter_by(is_active=True).all()])

@api.route("/warehouses", methods=["POST"])
@require_role("admin","manager")
def create_warehouse():
    data = request.get_json() or {}
    if not data.get("name"):
        return err("اسم المخزن مطلوب")
    wh = Warehouse(
        name     = data["name"],
        code     = data.get("code") or gen_ref("WH"),
        location = data.get("location"),
        capacity = data.get("capacity", 1000),
        type     = data.get("type","sub"),
    )
    db.session.add(wh)
    AuditService.log("create","warehouse",desc=f"مخزن: {wh.name}")
    db.session.commit()
    return created(wh.to_dict())

@api.route("/warehouses/<int:wid>", methods=["PUT"])
@require_role("admin","manager")
def update_warehouse(wid):
    wh   = Warehouse.query.get_or_404(wid)
    data = request.get_json() or {}
    for f in ["name","location","capacity","type","is_active"]:
        if f in data: setattr(wh,f,data[f])
    AuditService.log("update","warehouse",wid,f"تحديث: {wh.name}")
    db.session.commit()
    return ok(wh.to_dict())

# ══════════════════════════════════════════════════════════════
#  CATEGORIES
# ══════════════════════════════════════════════════════════════
@api.route("/categories", methods=["GET"])
@jwt_required()
def get_categories():
    return ok([c.to_dict() for c in Category.query.all()])

@api.route("/categories", methods=["POST"])
@require_role("admin","manager")
def create_category():
    data = request.get_json() or {}
    if not data.get("name"): return err("اسم الفئة مطلوب")
    if Category.query.filter_by(name=data["name"]).first():
        return err("الفئة موجودة مسبقاً")
    c = Category(name=data["name"],color=data.get("color","#3498db"),icon=data.get("icon","📦"))
    db.session.add(c); db.session.commit()
    return created(c.to_dict())

@api.route("/categories/<int:cid>", methods=["PUT"])
@require_role("admin","manager")
def update_category(cid):
    c    = Category.query.get_or_404(cid)
    data = request.get_json() or {}
    for f in ["name","color","icon"]:
        if f in data: setattr(c,f,data[f])
    db.session.commit()
    return ok(c.to_dict())

# ══════════════════════════════════════════════════════════════
#  ITEMS
# ══════════════════════════════════════════════════════════════
@api.route("/items", methods=["GET"])
@jwt_required()
def get_items():
    """
    List items
    ---
    parameters:
      - name: search
        in: query
        type: string
        required: false
      - name: category_id
        in: query
        type: integer
        required: false
      - name: status
        in: query
        type: string
        required: false
      - name: page
        in: query
        type: integer
        default: 1
      - name: per_page
        in: query
        type: integer
        default: 50
    responses:
      200: {description: Paginated items list}
    """
    result = ItemService.get_list(
        search      = request.args.get("search",""),
        category_id = request.args.get("category_id"),
        status      = request.args.get("status",""),
        page        = int(request.args.get("page",1)),
        per_page    = int(request.args.get("per_page",50)),
    )
    return ok({"items":result["items"],"total":result["total"],
               "page":result["page"],"pages":result["pages"]})

@api.route("/items/scan/<string:code>", methods=["GET"])
@jwt_required()
def scan_item(code):
    item = ItemService.scan(code)
    if not item: return not_found(f"لا يوجد صنف بالكود: {code}")
    return ok(item.to_dict())

@api.route("/items/<int:iid>", methods=["GET"])
@jwt_required()
def get_item(iid):
    item = Item.query.get_or_404(iid)
    d    = item.to_dict()
    d["last_movements"] = [m.to_dict() for m in
        StockMovement.query.filter_by(item_id=iid)
        .order_by(StockMovement.created_at.desc()).limit(10).all()]
    return ok(d)

@api.route("/items", methods=["POST"])
@jwt_required()
def create_item():
    uid  = int(get_jwt_identity())
    user = User.query.get(uid)
    if user.role == "viewer": return forbidden("ليس لديك صلاحية")
    data = request.get_json() or {}
    errs = validate(data, ["name","code"])
    if errs: return err("بيانات ناقصة", errors=errs)
    item, error = ItemService.create(data, uid)
    if error: return err(error)
    return created(item.to_dict())

@api.route("/items/<int:iid>", methods=["PUT"])
@jwt_required()
def update_item(iid):
    uid = int(get_jwt_identity())
    data = request.get_json() or {}
    item, error = ItemService.update(iid, data, uid)
    if error: return err(error)
    return ok(item.to_dict())

@api.route("/items/<int:iid>", methods=["DELETE"])
@require_role("admin")
def delete_item(iid):
    ok_flag, error = ItemService.delete(iid)
    if error: return err(error)
    return ok(message="تم حذف الصنف")

@api.route("/items/import", methods=["POST"])
@jwt_required()
@require_role("admin","manager")
def import_items():
    f = request.files.get("file")
    if not f: return err("الملف مطلوب")
    ext = f.filename.rsplit(".",1)[-1].lower() if "." in f.filename else ""
    imported = 0; errors = []
    rows = []
    if ext == "csv":
        content = f.read().decode("utf-8-sig").splitlines()
        reader = csv.DictReader(content)
        for row in reader: rows.append(row)
    elif ext in ("xlsx","xls"):
        import openpyxl
        wb = openpyxl.load_workbook(f)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
        for r in ws.iter_rows(min_row=2, values_only=True):
            rows.append(dict(zip(header, r)))
    else:
        return err("الرجاء رفع ملف CSV أو Excel")
    uid = int(get_jwt_identity())
    for i, row in enumerate(rows, 2):
        try:
            name = str(row.get("name","") or row.get("اسم الصنف","") or "").strip()
            code = str(row.get("code","") or row.get("الكود","") or "").strip()
            if not name or not code:
                errors.append(f"سطر {i}: الكود والاسم مطلوبان"); continue
            if Item.query.filter_by(code=code).first():
                errors.append(f"سطر {i}: الكود '{code}' موجود مسبقاً"); continue
            cat_name = str(row.get("category","") or row.get("الفئة","") or "").strip()
            cat = None
            if cat_name:
                cat = Category.query.filter_by(name=cat_name).first()
                if not cat:
                    cat = Category(name=cat_name); db.session.add(cat); db.session.flush()
            item = Item(
                code=code, name=name,
                barcode=str(row.get("barcode","") or row.get("الباركود","") or "") or None,
                category_id=cat.id if cat else None,
                unit=str(row.get("unit","") or row.get("الوحدة","") or "قطعة"),
                min_quantity=float(row.get("min_quantity","") or row.get("الحد الأدنى","") or 0),
                reorder_point=float(row.get("reorder_point","") or row.get("نقطة إعادة الطلب","") or 0),
                unit_price=float(row.get("unit_price","") or row.get("سعر الوحدة","") or 0),
                description=str(row.get("description","") or row.get("الوصف","") or ""),
            )
            db.session.add(item); imported += 1
        except Exception as e:
            errors.append(f"سطر {i}: {str(e)}")
    db.session.commit()
    AuditService.log("import","item",0,f"استيراد {imported} صنفاً عبر Excel/CSV")
    return ok({"imported": imported, "errors": errors[:50]}, f"تم استيراد {imported} صنفاً")

@api.route("/items/<int:iid>/barcode", methods=["GET"])
@jwt_required()
def item_barcode_img(iid):
    item = Item.query.get_or_404(iid)
    code = item.barcode or item.code
    try:
        import barcode as bc_lib
        from barcode.writer import ImageWriter
        code128 = bc_lib.get_barcode_class("code128")
        bc = code128(code, writer=ImageWriter())
        buf = io.BytesIO(); bc.write(buf); buf.seek(0)
        return send_file(buf, mimetype="image/png",
                         download_name=f"{item.code}_barcode.png")
    except ImportError:
        return err("مكتبة barcode غير مثبتة. قم بتشغيل: pip install python-barcode")

@api.route("/items/<int:iid>/barcode/label", methods=["GET"])
@jwt_required()
def item_barcode_label(iid):
    item = Item.query.get_or_404(iid)
    from reportlab.lib.pagesizes import label as lbl_size
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Image as RLImage, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    import tempfile, os
    code = item.barcode or item.code
    tmp = os.path.join(tempfile.gettempdir(), f"bcode_{item.id}.png")
    try:
        import barcode as bc_lib
        from barcode.writer import ImageWriter
        code128 = bc_lib.get_barcode_class("code128")
        bc = code128(code, writer=ImageWriter())
        with open(tmp, "wb") as f: bc.write(f)
    except ImportError:
        tmp = None
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=(100*mm, 60*mm),
                            rightMargin=5*mm, leftMargin=5*mm,
                            topMargin=5*mm, bottomMargin=5*mm)
    styles = getSampleStyleSheet()
    elems = [Paragraph(f"<b>{item.name}</b>", styles["Normal"]),
             Paragraph(f"الكود: {code} | {item.unit}", styles["Normal"])]
    if tmp and os.path.isfile(tmp):
        elems.append(Spacer(1, 3*mm))
        elems.append(RLImage(tmp, width=60*mm, height=20*mm))
    doc.build(elems); buf.seek(0)
    if tmp and os.path.isfile(tmp):
        try: os.remove(tmp)
        except: pass
    return send_file(buf, mimetype="application/pdf",
                     download_name=f"{item.code}_label.pdf")

@api.route("/items/bulk-labels", methods=["POST"])
@jwt_required()
def bulk_barcode_labels():
    data = request.get_json() or {}
    ids = data.get("item_ids", [])
    if not ids: return err("قائمة معرفات الأصناف مطلوبة")
    items = Item.query.filter(Item.id.in_(ids), Item.is_active==True).all()
    if not items: return err("لا توجد أصناف صالحة")
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Image as RLImage, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    import tempfile
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            rightMargin=10*mm, leftMargin=10*mm,
                            topMargin=10*mm, bottomMargin=10*mm)
    styles = getSampleStyleSheet()
    elems = []
    for item in items:
        code = item.barcode or item.code
        png = os.path.join(tempfile.gettempdir(), f"bc_{item.id}.png")
        try:
            import barcode as bc_lib
            from barcode.writer import ImageWriter
            code128 = bc_lib.get_barcode_class("code128")
            bc = code128(code, writer=ImageWriter())
            with open(png, "wb") as f: bc.write(f)
        except ImportError:
            png = None
        elems.append(Paragraph(f"<b>{item.name}</b>", styles["Normal"]))
        elems.append(Paragraph(f"الكود: {code} | {item.unit}", styles["Normal"]))
        if png and os.path.isfile(png):
            elems.append(RLImage(png, width=80*mm, height=25*mm))
        elems.append(Spacer(1, 5*mm))
    doc.build(elems); buf.seek(0)
    for item in items:
        p = os.path.join(tempfile.gettempdir(), f"bc_{item.id}.png")
        if os.path.isfile(p): os.remove(p)
    return send_file(buf, mimetype="application/pdf",
                     download_name=f"bulk_barcodes_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
                     as_attachment=True)

@api.route("/items/<int:iid>/image", methods=["POST"])
@jwt_required()
def upload_item_image(iid):
    item = Item.query.get_or_404(iid)
    f = request.files.get("file")
    if not f: return err("الملف مطلوب")
    ext = f.filename.rsplit(".",1)[-1].lower() if "." in f.filename else ""
    if ext not in {"jpg","jpeg","png","gif","webp"}:
        return err("نوع الملف غير مسموح: JPG, PNG, GIF, WEBP فقط")
    f.seek(0, os.SEEK_END); fsize = f.tell(); f.seek(0)
    if fsize > 5 * 1024 * 1024: return err("حجم الملف يتجاوز 5 ميجابايت")
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = f"item_img_{iid}_{ts}.{ext}"
    folder = os.path.join(current_app.root_path, "..", "uploads", "images")
    os.makedirs(folder, exist_ok=True)
    f.save(os.path.join(folder, safe_name))
    item.image_url = f"/uploads/images/{safe_name}"
    db.session.commit()
    AuditService.log("upload","item_image",iid,f"رفع صورة للصنف {item.name}")
    return ok({"image_url": item.image_url})

@api.route("/uploads/images/<path:filename>", methods=["GET"])
@jwt_required()
def serve_item_image(filename):
    folder = os.path.join(current_app.root_path, "..", "uploads", "images")
    return send_from_directory(folder, filename)

# ══════════════════════════════════════════════════════════════
#  STOCK MOVEMENTS
# ══════════════════════════════════════════════════════════════
@api.route("/movements", methods=["GET"])
@jwt_required()
def get_movements():
    filters = {
        "type":         request.args.get("type"),
        "warehouse_id": request.args.get("warehouse_id"),
        "item_id":      request.args.get("item_id"),
        "date_from":    request.args.get("date_from"),
        "date_to":      request.args.get("date_to"),
        "search":       request.args.get("search"),
    }
    result = MovementService.get_list(
        filters,
        page     = int(request.args.get("page",1)),
        per_page = int(request.args.get("per_page",30)),
    )
    return ok({"movements":[m.to_dict() for m in result["items"]],
               "total":result["total"],"page":result["page"],"pages":result["pages"]})

@api.route("/movements/<int:mid>/invoice", methods=["GET"])
@jwt_required()
def movement_invoice(mid):
    mov = StockMovement.query.get_or_404(mid)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    except ImportError:
        return err("reportlab غير مثبت")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=1.5*cm, leftMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    ts = ParagraphStyle("T", parent=styles["Title"], fontSize=16, alignment=TA_CENTER)
    elems = []
    company = "شركة النبلاء للمقاولات العامة والإنشاءات"
    elems.append(Paragraph(f"<b>{company}</b>", ts))
    elems.append(Paragraph(f"نظام إدارة المخازن — فاتورة صرف", styles["Normal"]))
    elems.append(Spacer(1, 0.5*cm))
    item_name = mov.item.name if mov.item else "—"
    wh_name = mov.warehouse.name if mov.warehouse else "—"
    uname = mov.user.name if mov.user else "—"
    info = [
        ["رقم الفاتورة:", f"INV-{mov.id:05d}"],
        ["التاريخ:", mov.created_at.strftime("%d/%m/%Y %H:%M") if mov.created_at else "—"],
        ["الصنف:", item_name],
        ["الكمية:", str(mov.quantity)],
        ["سعر الوحدة:", f"{mov.unit_price:,.2f} ر.س" if mov.unit_price else "—"],
        ["الإجمالي:", f"{mov.quantity * (mov.unit_price or 0):,.2f} ر.س"],
        ["المخزن:", wh_name],
        ["المرجع:", mov.ref_number or "—"],
        ["المنفذ:", uname],
        ["المشروع:", mov.project or "—"],
        ["المهندس:", mov.engineer_name or "—"],
    ]
    t = Table(info, colWidths=[4*cm, 10*cm])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0,0), (-1,-1), 10),
        ("ALIGN", (0,0), (0,-1), "RIGHT"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("TOPPADDING", (0,0), (-1,-1), 4),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
        ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#BDC3C7")),
        ("BACKGROUND", (0,0), (0,-1), colors.HexColor("#1B4F72")),
        ("TEXTCOLOR", (0,0), (0,-1), colors.white),
    ]))
    elems.append(t)
    if mov.notes:
        elems.append(Spacer(1, 0.3*cm))
        elems.append(Paragraph(f"<b>ملاحظات:</b> {mov.notes}", styles["Normal"]))
    doc.build(elems); buf.seek(0)
    return send_file(buf, mimetype="application/pdf",
                     download_name=f"invoice_{mov.id:05d}.pdf", as_attachment=True)

@api.route("/movements", methods=["POST"])
@jwt_required()
def create_movement():
    """
    Create stock movement
    ---
    parameters:
      - name: body
        in: body
        required: true
        schema:
          properties:
            item_id: {type: integer}
            warehouse_id: {type: integer}
            quantity: {type: number}
            type: {type: string, enum: [in, out, damage, return]}
            unit_price: {type: number}
    responses:
      201: {description: Movement created}
    """
    uid  = int(get_jwt_identity())
    user = User.query.get(uid)
    if user.role == "viewer": return forbidden("ليس لديك صلاحية تسجيل الحركات")
    data = request.get_json() or {}
    errs = validate(data, ["item_id","warehouse_id","quantity","type"])
    if errs: return err("بيانات ناقصة", errors=errs)
    mov, error = MovementService.create(data, uid)
    if error: return err(error)
    # low-stock email alert
    if data.get("type") in ("out","damage","transfer"):
        item = Item.query.get(data["item_id"])
        if item and item.get_total_stock() <= item.min_quantity:
            send_notif_email(f"⚠️ تنبيه نفاد: {item.name}",
                f"نفاذ مخزون الصنف: {item.name} ({item.code})\n"
                f"المتبقي: {item.get_total_stock()} {item.unit}\n"
                f"الحد الأدنى: {item.min_quantity} {item.unit}")
    return created(mov.to_dict())

# ══════════════════════════════════════════════════════════════
#  TRANSFERS
# ══════════════════════════════════════════════════════════════
@api.route("/transfers", methods=["GET"])
@jwt_required()
def get_transfers():
    status = request.args.get("status")
    return ok([t.to_dict() for t in TransferService.get_list(status)])

@api.route("/transfers", methods=["POST"])
@jwt_required()
def create_transfer():
    uid = int(get_jwt_identity())
    data = request.get_json() or {}
    errs = validate(data, ["item_id","from_warehouse_id","to_warehouse_id","quantity","reason"])
    if errs: return err("بيانات ناقصة", errors=errs)
    tr, error = TransferService.create(data, uid)
    if error: return err(error)
    item = Item.query.get(data["item_id"])
    to_wh = Warehouse.query.get(data["to_warehouse_id"])
    send_notif_email(f"🔄 تحويل جديد: {tr.ref_number}",
        f"تم إنشاء تحويل جديد:\nالصنف: {item.name if item else '—'}\n"
        f"الكمية: {data['quantity']}\nإلى: {to_wh.name if to_wh else '—'}\n"
        f"السبب: {data.get('reason','')}\nالطالب: {_uname(uid)}")
    return created(tr.to_dict())

def _process_chain_approval(entity, uid, resource_type):
    """Process multi-level approval. Returns (success, message, status)"""
    if not entity.approval_chain_id:
        return True, "", ""
    chain = ApprovalChain.query.get(entity.approval_chain_id)
    if not chain or not chain.is_active:
        return True, "", ""
    steps = sorted(chain.steps, key=lambda s: s.step_order)
    if not steps:
        return True, "", ""
    step_idx = entity.current_step or 0
    if step_idx >= len(steps):
        return True, "", ""
    step = steps[step_idx]
    user = User.query.get(uid)
    if user.role != step.role:
        return False, f"يجب أن يكون المعتمد بدور {step.role}", ""
    al = ApprovalLog(resource_type=resource_type, resource_id=entity.id,
                     approver_id=uid, status="approved", level=step_idx+1)
    db.session.add(al)
    if step.approval_type == "all" and step.role:
        approved = ApprovalLog.query.filter_by(
            resource_type=resource_type, resource_id=entity.id,
            level=step_idx+1, status="approved").count()
        total = User.query.filter_by(role=step.role).count()
        if approved < total:
            db.session.commit()
            return True, f"تمت الموافقة ({approved}/{total})", "step_same"
    entity.current_step = step_idx + 1
    if entity.current_step >= len(steps):
        return True, "", "final"
    db.session.commit()
    return True, f"تمت الموافقة على الخطوة {step_idx+1} من {len(steps)}", "step"

@api.route("/transfers/<int:tid>/approve", methods=["POST"])
@require_role("admin","manager")
def approve_transfer(tid):
    uid = int(get_jwt_identity())
    tr = Transfer.query.get_or_404(tid)
    if tr.status != "pending":
        return err("تمت معالجة هذا الطلب مسبقاً")
    is_ok, msg, final = _process_chain_approval(tr, uid, "transfer")
    if not is_ok: return err(msg)
    if final in ("step", "step_same"):
        db.session.commit()
        return ok(tr.to_dict(), msg)
    result, error = TransferService.approve(tid, uid)
    if error: return err(error)
    return ok(result.to_dict(), "تم اعتماد وتنفيذ التحويل")

@api.route("/transfers/<int:tid>/reject", methods=["POST"])
@require_role("admin","manager")
def reject_transfer(tid):
    uid = int(get_jwt_identity())
    data   = request.get_json() or {}
    tr, error = TransferService.reject(tid, uid, data.get("reason",""))
    if error: return err(error)
    return ok(tr.to_dict(), "تم رفض الطلب")

# ══════════════════════════════════════════════════════════════
#  STOCK ADJUSTMENT
# ══════════════════════════════════════════════════════════════
@api.route("/stock/adjust", methods=["POST"])
@jwt_required()
def adjust_stock():
    uid  = int(get_jwt_identity())
    user = User.query.get(uid)
    if user.role == "viewer": return forbidden("ليس لديك صلاحية")
    data = request.get_json() or {}
    errs = validate(data, ["item_id","warehouse_id","new_quantity","reason"])
    if errs: return err("بيانات ناقصة", errors=errs)
    mov, error = MovementService.adjust(data, uid)
    if error: return err(error)
    return ok(mov.to_dict(), "تمت التسوية")

# ══════════════════════════════════════════════════════════════
#  INVENTORY COUNTS
# ══════════════════════════════════════════════════════════════
@api.route("/counts", methods=["GET"])
@jwt_required()
def get_counts():
    return ok([c.to_dict(include_lines=False) for c in CountService.get_list()])

@api.route("/counts", methods=["POST"])
@jwt_required()
def create_count():
    uid = int(get_jwt_identity())
    data = request.get_json() or {}
    count, error = CountService.create(data, uid)
    if error: return err(error)
    return created(count.to_dict())

@api.route("/counts/<int:cid>", methods=["GET"])
@jwt_required()
def get_count(cid):
    count = InventoryCount.query.get_or_404(cid)
    return ok(count.to_dict())

@api.route("/counts/<int:cid>/lines/<int:lid>", methods=["PUT"])
@jwt_required()
def update_count_line(cid, lid):
    data = request.get_json() or {}
    line, error = CountService.update_line(
        cid, lid,
        data.get("actual_quantity"),
        data.get("notes",""),
    )
    if error: return err(error)
    return ok(line.to_dict())

@api.route("/counts/<int:cid>/complete", methods=["POST"])
@jwt_required()
def complete_count(cid):
    uid = int(get_jwt_identity())
    result, error = CountService.complete(cid, uid)
    if error: return err(error)
    return ok({"differences":result["differences"]}, "تم إغلاق الجرد وتسوية الفروقات")

@api.route("/schedule-counts", methods=["GET"])
@jwt_required()
def get_schedule_counts():
    return ok([s.to_dict() for s in ScheduledCount.query.all()])

@api.route("/schedule-counts", methods=["POST"])
@require_role("admin","manager")
def create_schedule_count():
    data = request.get_json() or {}
    sc = ScheduledCount(warehouse_id=data.get("warehouse_id"),
        frequency=data.get("frequency","monthly"),
        day_of_month=data.get("day_of_month",1),
        day_of_week=data.get("day_of_week",0))
    db.session.add(sc); db.session.commit()
    AuditService.log("create","schedule_count",sc.id,"جدولة جرد دوري")
    return created(sc.to_dict())

@api.route("/schedule-counts/<int:sid>", methods=["PUT"])
@require_role("admin","manager")
def update_schedule_count(sid):
    sc = ScheduledCount.query.get_or_404(sid)
    data = request.get_json() or {}
    for f in ["warehouse_id","frequency","day_of_month","day_of_week","is_active"]:
        if f in data: setattr(sc, f, data[f])
    db.session.commit()
    return ok(sc.to_dict())

@api.route("/schedule-counts/<int:sid>", methods=["DELETE"])
@require_role("admin")
def delete_schedule_count(sid):
    sc = ScheduledCount.query.get_or_404(sid)
    db.session.delete(sc); db.session.commit()
    return ok(message="تم حذف الجدولة")

# ══════════════════════════════════════════════════════════════
#  SUPPLIERS
# ══════════════════════════════════════════════════════════════
@api.route("/suppliers", methods=["GET"])
@jwt_required()
def get_suppliers():
    search = request.args.get("search","")
    return ok([s.to_dict() for s in SupplierService.get_list(search)])

@api.route("/suppliers", methods=["POST"])
@require_role("admin","manager")
def create_supplier():
    data = request.get_json() or {}
    sup, error = SupplierService.create(data)
    if error: return err(error)
    return created(sup.to_dict())

@api.route("/suppliers/<int:sid>", methods=["PUT"])
@require_role("admin","manager")
def update_supplier(sid):
    data = request.get_json() or {}
    sup, error = SupplierService.update(sid, data)
    if error: return err(error)
    return ok(sup.to_dict())

@api.route("/suppliers/<int:sid>", methods=["DELETE"])
@require_role("admin")
def delete_supplier(sid):
    s = Supplier.query.get_or_404(sid)
    s.is_active = False
    AuditService.log("delete","supplier",sid,f"حذف مورد: {s.name}")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  PROJECTS
# ══════════════════════════════════════════════════════════════
@api.route("/projects", methods=["GET"])
@jwt_required()
def get_projects():
    search = request.args.get("search","")
    q = Project.query
    if search:
        q = q.filter(Project.name.ilike(f"%{search}%"))
    return ok([p.to_dict() for p in q.order_by(Project.name).all()])

@api.route("/projects", methods=["POST"])
@require_role("admin","manager")
def create_project():
    data = request.get_json() or {}
    name = data.get("name","").strip()
    if not name: return err("اسم المشروع مطلوب")
    p = Project(name=name, code=data.get("code",""), description=data.get("description",""),
                status="active", budget=float(data.get("budget",0)),
                actual_cost=float(data.get("actual_cost",0)),
                completion_pct=float(data.get("completion_pct",0)),
                client=data.get("client",""), location=data.get("location",""))
    db.session.add(p)
    AuditService.log("create","project",p.id,f"إنشاء مشروع: {p.name}")
    db.session.commit()
    return created(p.to_dict())

@api.route("/projects/<int:pid>", methods=["PUT"])
@require_role("admin","manager")
def update_project(pid):
    p = Project.query.get_or_404(pid)
    data = request.get_json() or {}
    for f in ("name","code","description","status","client","location"):
        if f in data: setattr(p, f, data[f])
    for f in ("budget","actual_cost","completion_pct"):
        if f in data: setattr(p, f, float(data[f]))
    AuditService.log("edit","project",pid,f"تعديل مشروع: {p.name}")
    db.session.commit()
    return ok(p.to_dict())

@api.route("/projects/<int:pid>", methods=["DELETE"])
@require_role("admin")
def delete_project(pid):
    p = Project.query.get_or_404(pid)
    db.session.delete(p)
    AuditService.log("delete","project",pid,f"حذف مشروع: {p.name}")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  PROJECT INVOICES (مستخلصات)
# ══════════════════════════════════════════════════════════════
@api.route("/projects/<int:pid>/invoices", methods=["GET"])
@jwt_required()
def get_project_invoices(pid):
    p = Project.query.get_or_404(pid)
    invs = ProjectInvoice.query.filter_by(project_id=pid).order_by(ProjectInvoice.id.desc()).all()
    return ok([i.to_dict() for i in invs])

@api.route("/projects/<int:pid>/invoices", methods=["POST"])
@require_role("admin","manager")
def create_project_invoice(pid):
    p = Project.query.get_or_404(pid)
    data = request.get_json() or {}
    inv = ProjectInvoice(project_id=pid,
                         ref_number=data.get("ref_number",""),
                         amount=float(data.get("amount",0)),
                         description=data.get("description",""),
                         status=data.get("status","pending"))
    db.session.add(inv)
    AuditService.log("create","project_invoice",inv.id,f"إضافة مستخلص بمبلغ {inv.amount} لمشروع {p.name}")
    db.session.commit()
    return created(inv.to_dict())

@api.route("/projects/invoices/<int:iid>", methods=["PUT"])
@require_role("admin","manager")
def update_project_invoice(iid):
    inv = ProjectInvoice.query.get_or_404(iid)
    data = request.get_json() or {}
    for f in ("ref_number","description","status"):
        if f in data: setattr(inv, f, data[f])
    if "amount" in data:
        inv.amount = float(data["amount"])
    AuditService.log("edit","project_invoice",iid,f"تعديل مستخلص")
    db.session.commit()
    return ok(inv.to_dict())

@api.route("/projects/invoices/<int:iid>", methods=["DELETE"])
@require_role("admin")
def delete_project_invoice(iid):
    inv = ProjectInvoice.query.get_or_404(iid)
    pid = inv.project_id
    pname = inv.project.name
    db.session.delete(inv)
    AuditService.log("delete","project_invoice",iid,f"حذف مستخلص من مشروع {pname}")
    db.session.commit()
    return ok(message="تم حذف المستخلص")


@api.route("/stock/search", methods=["GET"])
@jwt_required()
def stock_search():
    q = request.args.get("q", "")
    return ok(StockInquiryService.search(q))

# ══════════════════════════════════════════════════════════════
#  USERS
# ══════════════════════════════════════════════════════════════
@api.route("/users", methods=["GET"])
@require_role("admin","manager")
def get_users():
    return ok([u.to_dict() for u in UserService.get_list()])

@api.route("/users", methods=["POST"])
@require_role("admin")
def create_user():
    data = request.get_json() or {}
    user, errors = UserService.create(data)
    if errors:
        msg = errors if isinstance(errors,str) else list(errors.values())[0]
        return err(msg, errors=errors if isinstance(errors,dict) else None)
    return created(user.to_dict())

@api.route("/users/<int:uid>", methods=["PUT"])
@require_role("admin")
def update_user(uid):
    current_uid = int(get_jwt_identity())
    data = request.get_json() or {}
    user, error = UserService.update(uid, data, current_uid)
    if error: return err(error)
    return ok(user.to_dict())

@api.route("/users/<int:uid>", methods=["DELETE"])
@require_role("admin")
def delete_user(uid):
    if uid == get_jwt_identity():
        return err("لا يمكن حذف حسابك الخاص")
    u = User.query.get_or_404(uid)
    u.is_active = False
    AuditService.log("delete","user",uid,f"حذف مستخدم: {u.name}")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  NOTIFICATIONS
# ══════════════════════════════════════════════════════════════
@api.route("/notifications", methods=["GET"])
@jwt_required()
def get_notifications():
    uid = int(get_jwt_identity())
    notifs = (Notification.query
              .filter_by(user_id=uid)
              .order_by(Notification.created_at.desc())
              .limit(50).all())
    unread = Notification.query.filter_by(user_id=uid, is_read=False).count()
    return ok({"notifications":[n.to_dict() for n in notifs],"unread":unread})

@api.route("/notifications/read-all", methods=["POST"])
@jwt_required()
def read_all_notifications():
    uid = int(get_jwt_identity())
    Notification.query.filter_by(user_id=uid, is_read=False).update({"is_read":True})
    db.session.commit()
    return ok(message="تم تحديد الكل كمقروء")

@api.route("/notifications/<int:nid>/read", methods=["POST"])
@jwt_required()
def read_notification(nid):
    n = Notification.query.get_or_404(nid)
    n.is_read = True
    db.session.commit()
    return ok()

# ══════════════════════════════════════════════════════════════
#  AUDIT LOGS
# ══════════════════════════════════════════════════════════════
@api.route("/audit-logs", methods=["GET"])
@require_role("admin","manager")
def get_audit_logs():
    page     = int(request.args.get("page",1))
    per_page = int(request.args.get("per_page",50))
    filters = {}
    for k in ["user_id","action","resource","date_from","date_to","search"]:
        v = request.args.get(k)
        if v: filters[k] = v
    result   = AuditService.get_logs(page, per_page, filters)
    return ok({"logs":[l.to_dict() for l in result["items"]],
               "total":result["total"],"pages":result["pages"]})

@api.route("/audit-logs/export", methods=["GET"])
@require_role("admin","manager")
def export_audit_logs():
    import csv
    fmt = request.args.get("format","csv")
    filters = {}
    for k in ["user_id","action","resource","date_from","date_to","search"]:
        v = request.args.get(k)
        if v: filters[k] = v
    result = AuditService.get_logs(1, 99999, filters)
    logs = [l.to_dict() for l in result["items"]]
    output = io.StringIO()
    output.write("\ufeff")
    w = csv.writer(output)
    w.writerow(["#","التاريخ","المستخدم","الإجراء","العنصر","الوصف","IP","البيانات القديمة","البيانات الجديدة"])
    for l in logs:
        w.writerow([l["id"],l["created_date"],l["user_name"],l["action"],
                    f"{l['resource']}#{l['resource_id']}" if l['resource_id'] else l['resource'],
                    l["description"] or "",l["ip_address"] or "",
                    l["old_data"] or "",l["new_data"] or ""])
    buf = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    fname = f"audit_logs_{datetime.datetime.now().strftime('%Y%m%d')}.csv"
    return send_file(buf, mimetype="text/csv;charset=utf-8",
                     download_name=fname, as_attachment=True)

# ══════════════════════════════════════════════════════════════
#  REPORTS
# ══════════════════════════════════════════════════════════════
@api.route("/reports/balance", methods=["GET"])
@jwt_required()
def report_balance():
    """
    Stock balance report
    ---
    responses:
      200: {description: All items with quantities per warehouse}
    """
    return ok(ReportService.balance())

@api.route("/reports/daily", methods=["GET"])
@jwt_required()
def report_daily():
    df = request.args.get("date_from") or request.args.get("date") or today_str()
    dt = request.args.get("date_to")
    data, error = ReportService.daily(df, dt)
    if error: return err(error)
    return ok(data)

@api.route("/reports/transfers", methods=["GET"])
@jwt_required()
def report_transfers():
    df = request.args.get("date_from")
    dt = request.args.get("date_to")
    return ok(ReportService.transfers(df, dt))

@api.route("/reports/budget", methods=["GET"])
@jwt_required()
def report_budget():
    return ok(ReportService.budget())

# ══════════════════════════════════════════════════════════════
#  EXPORT
# ══════════════════════════════════════════════════════════════
@api.route("/export/excel", methods=["GET"])
@jwt_required()
def export_excel():
    rtype  = request.args.get("type","balance")
    buf    = ReportService.export_excel(rtype)
    fname  = f"warehouse_{rtype}_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        download_name=fname, as_attachment=True,
    )

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT EXPORTS
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/export/<fmt>", methods=["GET"])
@jwt_required()
def procurement_export(fmt):
    from io import StringIO
    import csv
    rtype = request.args.get("type","pr")
    data = []
    headers = []
    fname_base = rtype
    if rtype == "pr":
        headers = ["رقم المرجع","التاريخ","الطالب","القسم","الأولوية","الحالة"]
        data = [[p.ref_number, str(p.created_at)[:10], p.requester.name if p.requester else "", p.department, p.priority, p.status] for p in PurchaseRequest.query.order_by(PurchaseRequest.created_at.desc()).all()]
    elif rtype == "rfq":
        headers = ["رقم المرجع","التاريخ","الحالة","عدد الموردين","عدد العروض"]
        data = [[r.ref_number, str(r.created_at)[:10], r.status, len(r.suppliers), len(r.quotations)] for r in RFQ.query.order_by(RFQ.created_at.desc()).all()]
    elif rtype == "po":
        headers = ["رقم المرجع","التاريخ","المورد","المستودع","الحالة","الإجمالي"]
        data = [[p.ref_number, str(p.created_at)[:10], p.supplier.name if p.supplier else "", p.warehouse.name if p.warehouse else "", p.status, p.total_amount] for p in PurchaseOrder.query.order_by(PurchaseOrder.created_at.desc()).all()]
    elif rtype == "returns":
        headers = ["رقم المرجع","التاريخ","المورد","السبب","الحالة"]
        data = [[r.ref_number, str(r.created_at)[:10], r.supplier.name if r.supplier else "", r.reason, r.status] for r in PurchaseReturn.query.order_by(PurchaseReturn.created_at.desc()).all()]
    elif rtype == "grn":
        headers = ["رقم المرجع","التاريخ","المورد","أمر الشراء","الملاحظات"]
        data = [[g.ref_number, str(g.created_at)[:10], g.supplier.name if g.supplier else "", g.po.ref_number if g.po else "", g.notes or ""] for g in GoodsReceipt.query.order_by(GoodsReceipt.created_at.desc()).all()]
    else:
        return not_found("نوع التقرير غير معروف")
    if fmt == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = fname_base
        ws.append(headers)
        for row in data: ws.append(row)
        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", download_name=f"procurement_{fname_base}.xlsx", as_attachment=True)
    elif fmt == "csv":
        buf = StringIO()
        w = csv.writer(buf); w.writerow(headers); w.writerows(data)
        b = io.BytesIO(buf.getvalue().encode("utf-8-sig"))
        return send_file(b, mimetype="text/csv", download_name=f"procurement_{fname_base}.csv", as_attachment=True)
    elif fmt == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.enums import TA_CENTER
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=1.5*cm, leftMargin=1.5*cm, topMargin=2*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        t_style = styles["Title"]
        elements = [Paragraph(f"تقرير المشتريات - {rtype}", t_style), Spacer(1, 0.5*cm)]
        col_w = [6*cm] + [4*cm]*(len(headers)-1) if len(headers) > 1 else [6*cm]
        tbl = Table([headers] + data, colWidths=col_w, repeatRows=1)
        ts = TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1B4F72")),
            ("TEXTCOLOR",(0,0),(-1,0),colors.white),
            ("FONTSIZE",(0,0),(-1,0),9),
            ("FONTSIZE",(0,1),(-1,-1),8),
            ("ALIGN",(0,0),(-1,-1),"CENTER"),
            ("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#BDC3C7")),
            ("TOPPADDING",(0,0),(-1,-1),4),
            ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ])
        tbl.setStyle(ts)
        elements.append(tbl)
        doc.build(elements); buf.seek(0)
        return send_file(buf, mimetype="application/pdf", download_name=f"procurement_{fname_base}.pdf", as_attachment=True)
    return not_found("نوع الملف غير معروف")

# ══════════════════════════════════════════════════════════════
#  INVENTORY VALUATION & ADVANCED REPORTS
# ══════════════════════════════════════════════════════════════
@api.route("/stock/valuation", methods=["GET"])
@jwt_required()
def stock_valuation():
    item_id = request.args.get("item_id", type=int)
    warehouse_id = request.args.get("warehouse_id", type=int)
    if item_id:
        return ok(InventoryLayer.get_valuation(item_id, warehouse_id))
    items = Item.query.filter_by(is_active=True).all()
    results = []
    grand_qty = 0; grand_val = 0.0
    for it in items:
        v = InventoryLayer.get_valuation(it.id, warehouse_id)
        if v["total_qty"] > 0:
            results.append({"item_id":it.id, "item_name":it.name, "item_code":it.code, "item_unit":it.unit, **v})
            grand_qty += v["total_qty"]; grand_val += v["total_value"]
    return ok({"layers": results, "total_items": len(results), "grand_total_qty": grand_qty, "grand_total_value": round(grand_val, 2)})

@api.route("/reports/aging", methods=["GET"])
@jwt_required()
def inventory_aging():
    items = Item.query.filter_by(is_active=True).all()
    now = datetime.datetime.utcnow()
    data = []
    for it in items:
        # age from FIFO layers
        layers = InventoryLayer.query.filter(InventoryLayer.item_id == it.id, InventoryLayer.quantity > 0).order_by(InventoryLayer.created_at).all()
        oldest = None; total_qty = 0; total_val = 0
        buckets = {"0-30":0, "31-60":0, "61-90":0, "91-180":0, "181-365":0, "365+":0}
        qty_buckets = {"0-30":0, "31-60":0, "61-90":0, "91-180":0, "181-365":0, "365+":0}
        for ly in layers:
            days = (now - ly.created_at).days
            q = ly.quantity
            total_qty += q; total_val += q * ly.unit_price
            if oldest is None or ly.created_at < oldest: oldest = ly.created_at
            if days <= 30: b = "0-30"
            elif days <= 60: b = "31-60"
            elif days <= 90: b = "61-90"
            elif days <= 180: b = "91-180"
            elif days <= 365: b = "181-365"
            else: b = "365+"
            buckets[b] = round(buckets[b] + q * ly.unit_price, 2)
            qty_buckets[b] += q
        if total_qty > 0:
            data.append({
                "item_id": it.id, "item_name": it.name, "item_code": it.code,
                "unit": it.unit, "total_qty": total_qty, "total_value": round(total_val, 2),
                "oldest_days": (now - oldest).days if oldest else 0,
                "buckets": buckets, "qty_buckets": qty_buckets,
            })
    data.sort(key=lambda x: x["oldest_days"], reverse=True)
    return ok({"items": data, "total_items": len(data)})

@api.route("/stock/cleanup-layers", methods=["POST"])
@jwt_required()
@require_role("admin","manager")
def cleanup_layers():
    deleted = InventoryLayer.query.filter(InventoryLayer.quantity == 0).delete()
    db.session.commit()
    AuditService.log("cleanup","inventory_layer",0,f"حذف {deleted} FIFO layer(s) صفر الكمية")
    return ok({"deleted": deleted}, f"تم حذف {deleted} طبقة")

# ══════════════════════════════════════════════════════════════
#  LOTS
# ══════════════════════════════════════════════════════════════
@api.route("/lots", methods=["GET"])
@jwt_required()
def get_lots():
    item_id = request.args.get("item_id", type=int)
    status  = request.args.get("status")
    q = Lot.query
    if item_id: q = q.filter_by(item_id=item_id)
    if status:  q = q.filter_by(status=status)
    return ok([l.to_dict() for l in q.order_by(Lot.created_at.desc()).all()])

@api.route("/lots", methods=["POST"])
@jwt_required()
@require_role("admin","manager")
def create_lot():
    data = request.get_json() or {}
    item_id = data.get("item_id")
    lot_num = data.get("lot_number","").strip()
    if not item_id or not lot_num: return err("الصنف ورقم الدفعة مطلوبان")
    existing = Lot.query.filter_by(item_id=item_id, lot_number=lot_num).first()
    if existing: return err("رقم الدفعة موجود مسبقاً لهذا الصنف")
    lot = Lot(item_id=item_id, lot_number=lot_num,
              expiry_date=datetime.datetime.strptime(data["expiry_date"],"%Y-%m-%d") if data.get("expiry_date") else None)
    db.session.add(lot); db.session.commit()
    AuditService.log("create","lot",lot.id,f"إضافة دفعة {lot_num}")
    return created(lot.to_dict())

@api.route("/lots/<int:lid>", methods=["PUT"])
@jwt_required()
@require_role("admin","manager")
def update_lot(lid):
    lot = Lot.query.get_or_404(lid)
    data = request.get_json() or {}
    for f in ["lot_number","status"]:
        if f in data: setattr(lot, f, data[f])
    if "expiry_date" in data:
        lot.expiry_date = datetime.datetime.strptime(data["expiry_date"],"%Y-%m-%d") if data["expiry_date"] else None
    db.session.commit()
    return ok(lot.to_dict())

# ══════════════════════════════════════════════════════════════
#  UNIT CONVERSIONS
# ══════════════════════════════════════════════════════════════
@api.route("/units", methods=["GET"])
@jwt_required()
def get_units():
    return ok([u.to_dict() for u in UnitConversion.query.all()])

@api.route("/units/convert", methods=["GET"])
@jwt_required()
def convert_unit():
    from_unit = request.args.get("from","")
    to_unit   = request.args.get("to","")
    qty       = request.args.get("qty", 1, type=float)
    if from_unit == to_unit:
        return ok({"from":from_unit,"to":to_unit,"qty":qty,"result":qty,"factor":1})
    conv = UnitConversion.query.filter_by(from_unit=from_unit, to_unit=to_unit).first()
    if not conv:
        # try reverse
        rev = UnitConversion.query.filter_by(from_unit=to_unit, to_unit=from_unit).first()
        if rev:
            return ok({"from":from_unit,"to":to_unit,"qty":qty,"result":round(qty / rev.factor, 4),"factor":round(1/rev.factor, 4)})
        return err("لا يوجد تحويل بين هاتين الوحدتين")
    return ok({"from":from_unit,"to":to_unit,"qty":qty,"result":round(qty * conv.factor, 4),"factor":conv.factor})

@api.route("/units", methods=["POST"])
@jwt_required()
@require_role("admin")
def create_unit():
    data = request.get_json() or {}
    if not data.get("from_unit") or not data.get("to_unit") or not data.get("factor"):
        return err("جميع الحقول مطلوبة")
    existing = UnitConversion.query.filter_by(from_unit=data["from_unit"], to_unit=data["to_unit"]).first()
    if existing: return err("التحويل موجود مسبقاً")
    uc = UnitConversion(from_unit=data["from_unit"], to_unit=data["to_unit"], factor=data["factor"])
    db.session.add(uc); db.session.commit()
    return created(uc.to_dict())

@api.route("/units/<int:uid>", methods=["DELETE"])
@jwt_required()
@require_role("admin")
def delete_unit(uid):
    uc = UnitConversion.query.get_or_404(uid)
    db.session.delete(uc); db.session.commit()
    return ok(message="تم حذف التحويل")

# ══════════════════════════════════════════════════════════════
#  APPROVAL CHAINS
# ══════════════════════════════════════════════════════════════
@api.route("/approval/chains", methods=["GET"])
@jwt_required()
def get_approval_chains():
    target_type = request.args.get("target_type")
    q = ApprovalChain.query
    if target_type: q = q.filter_by(target_type=target_type)
    return ok([c.to_dict() for c in q.order_by(ApprovalChain.created_at.desc()).all()])

@api.route("/approval/chains", methods=["POST"])
@jwt_required()
@require_role("admin")
def create_approval_chain():
    data = request.get_json() or {}
    if not data.get("name") or not data.get("target_type"):
        return err("الاسم ونوع الهدف مطلوبان")
    chain = ApprovalChain(name=data["name"], target_type=data["target_type"])
    db.session.add(chain); db.session.flush()
    for i, s in enumerate(data.get("steps", [])):
        step = ApprovalStep(chain_id=chain.id, step_order=i+1,
            role=s.get("role","manager"), user_id=s.get("user_id"),
            approval_type=s.get("approval_type","any"))
        db.session.add(step)
    db.session.commit()
    AuditService.log("create","approval_chain",chain.id,f"إنشاء سلسلة اعتماد: {chain.name}")
    return created(chain.to_dict())

@api.route("/approval/chains/<int:cid>", methods=["PUT"])
@jwt_required()
@require_role("admin")
def update_approval_chain(cid):
    chain = ApprovalChain.query.get_or_404(cid)
    data = request.get_json() or {}
    for f in ["name","target_type","is_active"]:
        if f in data: setattr(chain, f, data[f])
    if "steps" in data:
        ApprovalStep.query.filter_by(chain_id=cid).delete()
        for i, s in enumerate(data["steps"]):
            step = ApprovalStep(chain_id=cid, step_order=i+1,
                role=s.get("role","manager"), user_id=s.get("user_id"),
                approval_type=s.get("approval_type","any"))
            db.session.add(step)
    db.session.commit()
    return ok(chain.to_dict())

@api.route("/approval/chains/<int:cid>", methods=["DELETE"])
@jwt_required()
@require_role("admin")
def delete_approval_chain(cid):
    chain = ApprovalChain.query.get_or_404(cid)
    db.session.delete(chain); db.session.commit()
    return ok(message="تم حذف سلسلة الاعتماد")

@api.route("/reports/consumption", methods=["GET"])
@jwt_required()
def consumption_report():
    days = int(request.args.get("days", 30))
    since = datetime.datetime.utcnow() - datetime.timedelta(days=days)
    limit = int(request.args.get("limit", 20))
    order = request.args.get("order", "desc")
    rows = db.session.query(
        StockMovement.item_id, Item.name, Item.code, Item.unit,
        func.sum(StockMovement.quantity).label("total_qty"),
        func.sum(StockMovement.quantity * StockMovement.unit_price).label("total_val"),
        func.count(StockMovement.id).label("mov_count")
    ).join(Item, StockMovement.item_id == Item.id
    ).filter(StockMovement.type.in_(["out","transfer","damage"]),
             StockMovement.created_at >= since
    ).group_by(StockMovement.item_id
    ).order_by(func.sum(StockMovement.quantity).desc() if order=="desc" else func.sum(StockMovement.quantity).asc()
    ).limit(limit).all()
    data = [{"item_id":r[0],"item_name":r[1],"item_code":r[2],"item_unit":r[3],
             "total_qty":float(r[4]),"total_value":round(float(r[5] or 0),2),"movements":r[6],
             "avg_per_day":round(float(r[4])/max(days,1),2)} for r in rows]
    return ok({"items":data, "period_days":days, "order":order, "total":len(data)})

@api.route("/reports/fast-slow", methods=["GET"])
@jwt_required()
def fast_slow_report():
    days = int(request.args.get("days", 90))
    since = datetime.datetime.utcnow() - datetime.timedelta(days=days)
    limit = int(request.args.get("limit", 10))
    out_q = db.session.query(StockMovement.item_id, func.sum(StockMovement.quantity).label("qty")).filter(
        StockMovement.type.in_(["out","transfer","damage"]), StockMovement.created_at >= since
    ).group_by(StockMovement.item_id).subquery()
    items = db.session.query(Item.id, Item.name, Item.code, Item.unit, Item.unit_price,
                             func.coalesce(out_q.c.qty, 0).label("consumed"),
                             func.coalesce(func.sum(Stock.quantity), 0).label("current_stock")
    ).outerjoin(out_q, Item.id == out_q.c.item_id
    ).outerjoin(Stock, Stock.item_id == Item.id
    ).group_by(Item.id).all()
    fast = sorted(items, key=lambda x: float(x.consumed or 0), reverse=True)[:limit]
    slow = sorted([it for it in items if float(it.consumed or 0) <= 0], key=lambda x: float(x.current_stock or 0), reverse=True)[:limit]
    def fmt(i): return {"item_id":i[0],"item_name":i[1],"item_code":i[2],"item_unit":i[3],"unit_price":float(i[4] or 0),"consumed":float(i[5] or 0),"current_stock":float(i[6] or 0)}
    return ok({"fast_moving":[fmt(i) for i in fast],"slow_moving":[fmt(i) for i in slow],"period_days":days})

@api.route("/reports/supplier-performance", methods=["GET"])
@jwt_required()
def supplier_performance():
    rows = db.session.query(
        Supplier.id, Supplier.name,
        func.count(PurchaseOrder.id).label("po_count"),
        func.sum(PurchaseOrder.total_amount).label("total_amount"),
        func.count(GoodsReceipt.id).label("grn_count"),
    ).outerjoin(PurchaseOrder, PurchaseOrder.supplier_id == Supplier.id
    ).outerjoin(GoodsReceipt, GoodsReceipt.po_id == PurchaseOrder.id
    ).group_by(Supplier.id).all()
    data = [{"supplier_id":r[0],"supplier_name":r[1],"po_count":r[2],"total_amount":float(r[3] or 0),
             "grn_count":r[4]} for r in rows]
    evals = SupplierEvaluation.query.with_entities(
        SupplierEvaluation.supplier_id,
        func.avg(SupplierEvaluation.quality).label("avg_quality"),
        func.avg(SupplierEvaluation.delivery).label("avg_delivery"),
        func.avg(SupplierEvaluation.price).label("avg_price")
    ).group_by(SupplierEvaluation.supplier_id).all()
    eval_map = {e[0]:{"quality":round(float(e[1] or 0),1),"delivery":round(float(e[2] or 0),1),"price":round(float(e[3] or 0),1)} for e in evals}
    for d in data:
        d["evaluations"] = eval_map.get(d["supplier_id"], {})
    return ok(data)

@api.route("/reports/abc", methods=["GET"])
@jwt_required()
def abc_analysis():
    items = Item.query.filter_by(is_active=True).all()
    values = []
    for it in items:
        total_stock = it.get_total_stock()
        total_value = total_stock * it.unit_price
        values.append({"item_id":it.id,"item_name":it.name,"item_code":it.code,
                       "category":it.category.name if it.category else "",
                       "unit":it.unit,"total_stock":total_stock,
                       "unit_price":float(it.unit_price),
                       "total_value":round(total_value,2)})
    values.sort(key=lambda x: x["total_value"], reverse=True)
    grand_total = sum(v["total_value"] for v in values) or 1
    cumulative = 0
    for v in values:
        cumulative += v["total_value"]
        v["pct"] = round(v["total_value"] / grand_total * 100, 2)
        v["cumulative_pct"] = round(cumulative / grand_total * 100, 2)
        v["class"] = "A" if v["cumulative_pct"] <= 80 else "B" if v["cumulative_pct"] <= 95 else "C"
    a = [v for v in values if v["class"]=="A"]
    b = [v for v in values if v["class"]=="B"]
    c = [v for v in values if v["class"]=="C"]
    return ok({
        "items": values,
        "summary": {
            "grand_total": round(grand_total, 2),
            "A": {"count":len(a),"value":round(sum(v["total_value"] for v in a),2)},
            "B": {"count":len(b),"value":round(sum(v["total_value"] for v in b),2)},
            "C": {"count":len(c),"value":round(sum(v["total_value"] for v in c),2)},
        }
    })

# ══════════════════════════════════════════════════════════════
#  FILE ATTACHMENTS — PROCUREMENT
# ══════════════════════════════════════════════════════════════
@api.route("/attachments", methods=["POST"])
@jwt_required()
def upload_attachment():
    if "file" not in request.files:
        return err("الملف مطلوب")
    f = request.files["file"]
    if not f.filename: return err("اسم الملف مطلوب")
    ext = f.filename.rsplit(".",1)[-1].lower() if "." in f.filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        return err("نوع الملف غير مسموح: يرجى رفع PDF, صور, مستندات Office أو CSV/ZIP فقط")
    f.seek(0, os.SEEK_END)
    fsize = f.tell()
    f.seek(0)
    if fsize > MAX_FILE_SIZE:
        return err("حجم الملف يتجاوز 16 ميجابايت")
    ref_type = request.form.get("ref_type", "item")
    ref_id = request.form.get("ref_id", type=int)
    upload_dir = os.path.join(current_app.root_path, "..", "uploads")
    os.makedirs(upload_dir, exist_ok=True)
    safe_name = f"{ref_type}_{ref_id}_{int(datetime.datetime.utcnow().timestamp())}_{secure_filename(f.filename)}"
    path = os.path.join(upload_dir, safe_name)
    f.save(path)
    size = os.path.getsize(path)
    att = ItemAttachment(
        item_id=ref_id,
        filename=safe_name, original_name=f.filename,
        file_type=ref_type, file_size=size, notes=request.form.get("notes",""),
        uploaded_by=g.current_user.id,
    )
    db.session.add(att)
    AuditService.log("create","attachment",att.id,f"رفع ملف: {f.filename}")
    db.session.commit()
    return created(att.to_dict(), "تم رفع الملف")

@api.route("/attachments", methods=["GET"])
@jwt_required()
def list_attachments():
    ref_type = request.args.get("ref_type")
    ref_id = request.args.get("ref_id", type=int)
    q = ItemAttachment.query
    if ref_type: q = q.filter(ItemAttachment.file_type == ref_type)
    if ref_id: q = q.filter(ItemAttachment.item_id == ref_id)
    return ok([a.to_dict() for a in q.order_by(ItemAttachment.created_at.desc()).all()])

@api.route("/attachments/<int:aid>/download", methods=["GET"])
@jwt_required()
def download_attachment(aid):
    att = ItemAttachment.query.get_or_404(aid)
    upload_dir = os.path.join(current_app.root_path, "..", "uploads")
    return send_from_directory(upload_dir, att.filename, as_attachment=True, download_name=att.original_name)

@api.route("/attachments/<int:aid>", methods=["DELETE"])
@require_role("admin","manager")
def delete_attachment(aid):
    att = ItemAttachment.query.get_or_404(aid)
    path = os.path.join(current_app.root_path, "..", "uploads", att.filename)
    if os.path.isfile(path): os.remove(path)
    db.session.delete(att)
    AuditService.log("delete","attachment",aid,f"حذف ملف: {att.original_name}")
    db.session.commit()
    return ok(message="تم حذف الملف")

# ══════════════════════════════════════════════════════════════
#  CYCLE COUNT — VARIANCE REPORT & AUTO ADJUST
# ══════════════════════════════════════════════════════════════
@api.route("/inventory-counts/<int:cid>/variance", methods=["GET"])
@jwt_required()
def count_variance(cid):
    ic = InventoryCount.query.get_or_404(cid)
    return ok({"count": ic.to_dict(include_lines=True),
               "variance_lines": [l.to_dict() for l in ic.lines if l.difference is not None and l.difference != 0],
               "total_variance": sum(abs(l.difference or 0) for l in ic.lines if l.difference is not None)})

@api.route("/inventory-counts/<int:cid>/auto-adjust", methods=["POST"])
@require_role("admin","manager")
def auto_adjust_count(cid):
    ic = InventoryCount.query.get_or_404(cid)
    if ic.status != "completed": return err("يجب إكمال الجرد أولاً")
    adjustments = 0
    for line in ic.lines:
        if line.difference is not None and line.difference != 0:
            stk = Stock.get_or_create(line.item_id, ic.warehouse_id)
            stk.quantity = max(0, stk.quantity + line.difference)
            StockMovement(ref_number=gen_ref("ADJ"), type="adjustment",
                item_id=line.item_id, warehouse_id=ic.warehouse_id,
                quantity=abs(line.difference), unit_price=line.item.unit_price if line.item else 0,
                user_id=g.current_user.id, notes=f"تسوية جرد #{ic.ref_number}: فرق {line.difference:+.2f}")
            adjustments += 1
    AuditService.log("adjust","inventory_count",cid,f"تسوية جرد: {adjustments} صنف")
    db.session.commit()
    return ok(message=f"تم تسوية {adjustments} صنف")

# ══════════════════════════════════════════════════════════════
#  GENERATE PO FROM QUOTATION
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/quotations/<int:qid>/po", methods=["POST"])
@require_role("admin","manager")
def generate_po_from_quotation(qid):
    qt = Quotation.query.get_or_404(qid)
    if qt.status != "accepted": return err("يجب قبول عرض السعر أولاً")
    po = PurchaseOrder(ref_number=gen_ref("PO"), quotation_id=qt.id,
        supplier_id=qt.supplier_id, rfq_id=qt.rfq_id,
        warehouse_id=request.get_json(silent=True) and request.get_json(silent=True).get("warehouse_id") or None,
        created_by=g.current_user.id, status="draft",
        notes=qt.notes, total_amount=qt.total)
    db.session.add(po); db.session.flush()
    for qi in qt.items:
        db.session.add(POItem(po_id=po.id, item_name=qi.item_name,
            quantity=qi.quantity, unit_price=qi.unit_price,
            total=qi.total, description=qi.description))
    AuditService.log("create","purchase_order",po.id,f"إنشاء أمر شراء من عرض سعر: {po.ref_number}")
    db.session.commit()
    return created(po.to_dict(), "تم إنشاء أمر الشراء من عرض السعر")

# ══════════════════════════════════════════════════════════════
#  MISSING DETAIL ENDPOINTS
# ══════════════════════════════════════════════════════════════
@api.route("/warehouses/<int:wid>", methods=["GET"])
@jwt_required()
def get_warehouse(wid):
    return ok(Warehouse.query.get_or_404(wid).to_dict())

@api.route("/suppliers/<int:sid>", methods=["GET"])
@jwt_required()
def get_supplier(sid):
    return ok(Supplier.query.get_or_404(sid).to_dict())

@api.route("/users/<int:uid>", methods=["GET"])
@require_role("admin","manager")
def get_user(uid):
    return ok(User.query.get_or_404(uid).to_dict())

@api.route("/transfers/<int:tid>", methods=["GET"])
@jwt_required()
def get_transfer(tid):
    return ok(Transfer.query.get_or_404(tid).to_dict())

@api.route("/procurement/evaluations/<int:eid>", methods=["GET"])
@jwt_required()
def get_evaluation(eid):
    return ok(SupplierEvaluation.query.get_or_404(eid).to_dict())

@api.route("/procurement/grn/<int:gid>", methods=["PUT"])
@require_role("admin","manager")
def update_grn(gid):
    grn = GoodsReceipt.query.get_or_404(gid)
    data = request.get_json() or {}
    for k in ("notes",): setattr(grn, k, data.get(k, getattr(grn, k)))
    AuditService.log("update","goods_receipt",gid,f"تحديث إذن استلام")
    db.session.commit()
    return ok(grn.to_dict())

@api.route("/procurement/returns/<int:rid>", methods=["PUT"])
@require_role("admin","manager")
def update_return(rid):
    pr = PurchaseReturn.query.get_or_404(rid)
    data = request.get_json() or {}
    for k in ("notes","reason"): setattr(pr, k, data.get(k, getattr(pr, k)))
    AuditService.log("update","purchase_return",rid,f"تحديث مرتجع")
    db.session.commit()
    return ok(pr.to_dict())

# ══════════════════════════════════════════════════════════════
#  BACKUP
# ══════════════════════════════════════════════════════════════
@api.route("/backup", methods=["GET"])
@require_role("admin")
def backup():
    data = {
        "backup_time": datetime.datetime.now().isoformat(),
        "version":     "2.0",
        "warehouses":  [w.to_dict() for w in Warehouse.query.all()],
        "categories":  [c.to_dict() for c in Category.query.all()],
        "suppliers":   [s.to_dict() for s in Supplier.query.all()],
        "items":       [i.to_dict() for i in Item.query.all()],
        "movements":   [m.to_dict() for m in
                        StockMovement.query.order_by(StockMovement.created_at).all()],
    }
    buf  = io.BytesIO(json.dumps(data,ensure_ascii=False,indent=2).encode("utf-8"))
    buf.seek(0)
    fname= f"wms_backup_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    return send_file(buf, mimetype="application/json",
                     download_name=fname, as_attachment=True)

# ══════════════════════════════════════════════════════════════
#  REGISTER  — POST /api/auth/register
# ══════════════════════════════════════════════════════════════
@api.route("/auth/register", methods=["POST"])
@limiter.limit("3 per hour")
def register():
    data = request.get_json() or {}
    errs = validate(data, ["name", "username", "password"])
    if errs:
        return err("بيانات ناقصة", errors=errs)
    if len(data.get("password","")) < 6:
        return err("كلمة المرور 6 أحرف على الأقل")
    if User.query.filter_by(username=data["username"]).first():
        return err("اسم المستخدم مستخدم مسبقاً")
    if data.get("email") and User.query.filter_by(email=data["email"]).first():
        return err("البريد الإلكتروني مستخدم مسبقاً")
    u = User(name=data["name"], username=data["username"],
             email=data.get("email"), role="viewer")
    u.set_password(data["password"])
    db.session.add(u)
    AuditService.log("create","user",None,f"تسجيل جديد: {u.name}")
    db.session.commit()
    access  = create_access_token(identity=str(u.id))
    refresh = create_refresh_token(identity=str(u.id))
    return created({"access_token":access,"refresh_token":refresh,
                    "user":u.to_dict(),"permissions":u.get_permissions()},
                   "تم التسجيل بنجاح")


# ══════════════════════════════════════════════════════════════
#  CSV EXPORT  — GET /api/export/csv?type=balance|movements
# ══════════════════════════════════════════════════════════════
@api.route("/export/csv", methods=["GET"])
@jwt_required()
def export_csv():
    import csv
    rtype  = request.args.get("type","balance")
    output = io.StringIO()
    output.write("\ufeff")   # BOM for Arabic Excel
    writer = csv.writer(output)

    if rtype == "balance":
        whs = Warehouse.query.filter_by(is_active=True).all()
        writer.writerow(["الكود","الصنف","الفئة","الوحدة"] +
                        [w.name for w in whs] + ["الإجمالي","القيمة (ر.س)","الحالة"])
        for item in Item.query.filter_by(is_active=True).order_by(Item.name).all():
            stocks = {s.warehouse_id: s.quantity for s in item.stocks}
            total  = sum(stocks.values())
            writer.writerow(
                [item.code, item.name,
                 item.category.name if item.category else "", item.unit] +
                [stocks.get(w.id, 0) for w in whs] +
                [total, round(total * item.unit_price, 2),
                 {"critical":"حرج","warning":"تحذير","ok":"جيد"}[item.get_status()]]
            )
    elif rtype == "movements":
        writer.writerow(["#","التاريخ","الوقت","الصنف","النوع",
                         "الكمية","الوحدة","المخزن","المنفذ","المرجع","المشروع","المهندس","ملاحظات"])
        tmap = {"in":"وارد","out":"صرف","transfer":"تحويل","adjustment":"تسوية",
                "return":"مرتجع","damage":"هالك"}
        for m in (StockMovement.query
                   .order_by(StockMovement.created_at.desc()).limit(5000).all()):
            writer.writerow([
                m.id, m.created_at.strftime("%d/%m/%Y"), m.created_at.strftime("%H:%M"),
                m.item.name if m.item else "", tmap.get(m.type,""),
                m.quantity, m.item.unit if m.item else "",
                m.warehouse.name if m.warehouse else "",
                m.user.name    if m.user    else "",
                m.ref_number or "",
                m.project or "",
                m.engineer_name or "",
                m.notes or "",
            ])
    elif rtype == "transfers":
        writer.writerow(["#","رقم المرجع","الصنف","الكمية","الوحدة",
                         "من مخزن","إلى مخزن","الطالب","التاريخ","السبب","الحالة"])
        for t in Transfer.query.order_by(Transfer.created_at.desc()).all():
            writer.writerow([
                t.id, t.ref_number or "",
                t.item.name if t.item else "", t.quantity,
                t.item.unit if t.item else "",
                t.from_wh.name if t.from_wh else "",
                t.to_wh.name if t.to_wh else "",
                t.requester.name if t.requester else "",
                t.created_at.strftime("%d/%m/%Y") if t.created_at else "",
                t.reason or "",
                t.STATUS_LABELS.get(t.status,t.status),
            ])
    elif rtype == "diffs":
        writer.writerow(["#","الصنف","كود الصنف","المخزن","كمية النظام",
                         "الكمية الفعلية","الفرق","ملاحظات","تاريخ الجرد"])
        diffs = (db.session.query(InventoryCountLine)
                 .join(InventoryCount)
                 .filter(
                     InventoryCountLine.actual_quantity.isnot(None),
                     InventoryCountLine.actual_quantity != InventoryCountLine.system_quantity,
                     InventoryCount.status == "completed"
                 )
                 .order_by(InventoryCount.completed_at.desc())
                 .all())
        for l in diffs:
            writer.writerow([
                l.id,
                l.item.name if l.item else "",
                l.item.code if l.item else "",
                l.count.warehouse.name if l.count and l.count.warehouse else "",
                l.system_quantity, l.actual_quantity,
                l.difference, l.notes or "",
                l.count.completed_at.strftime("%d/%m/%Y") if l.count and l.count.completed_at else "",
            ])
    elif rtype == "budget":
        writer.writerow(["المشروع","الكود","العميل","الميزانية","المنصرف","المتبقي","الاستخدام %","الحالة"])
        data = ReportService.budget()
        for p in data["projects"]:
            writer.writerow([p["name"],p["code"],p["client"],
                            p["budget"],p["spent"],p["remaining"],
                            f"{p['usage_pct']}%",p["status_label"]])
    elif rtype == "consumption":
        days = int(request.args.get("days",30))
        since = datetime.datetime.utcnow() - datetime.timedelta(days=days)
        limit = int(request.args.get("limit",50))
        writer.writerow(["الصنف","الكود","الكمية المستهلكة","القيمة","عدد الحركات","متوسط اليوم"])
        rows = db.session.query(
            StockMovement.item_id, Item.name, Item.code,
            func.sum(StockMovement.quantity).label("total_qty"),
            func.sum(StockMovement.quantity * StockMovement.unit_price).label("total_val"),
            func.count(StockMovement.id).label("mov_count")
        ).join(Item, StockMovement.item_id == Item.id
        ).filter(StockMovement.type.in_(["out","transfer","damage"]),
                 StockMovement.created_at >= since
        ).group_by(StockMovement.item_id
        ).order_by(func.sum(StockMovement.quantity).desc()
        ).limit(limit).all()
        for r in rows:
            writer.writerow([r[1],r[2],r[3],round(float(r[4] or 0),2),r[5],round(float(r[3])/max(days,1),2)])
    elif rtype == "fastslow":
        days = int(request.args.get("days",90))
        since = datetime.datetime.utcnow() - datetime.timedelta(days=days)
        limit = int(request.args.get("limit",50))
        out_q = db.session.query(StockMovement.item_id, func.sum(StockMovement.quantity).label("qty")).filter(
            StockMovement.type.in_(["out","transfer","damage"]), StockMovement.created_at >= since
        ).group_by(StockMovement.item_id).subquery()
        items = db.session.query(Item.id, Item.name, Item.code, Item.unit,
                                 func.coalesce(out_q.c.qty, 0).label("consumed"),
                                 func.coalesce(func.sum(Stock.quantity), 0).label("current_stock")
        ).outerjoin(out_q, Item.id == out_q.c.item_id
        ).outerjoin(Stock, Stock.item_id == Item.id
        ).group_by(Item.id).all()
        writer.writerow(["النوع","الصنف","الكمية المستهلكة","المخزون الحالي"])
        fast = sorted(items, key=lambda x: float(x.consumed or 0), reverse=True)[:limit]
        slow = sorted([it for it in items if float(it.consumed or 0) <= 0], key=lambda x: float(x.current_stock or 0), reverse=True)[:limit]
        for i in fast:
            writer.writerow(["سريع الحركة",i[1],i[4],i[5]])
        for i in slow:
            writer.writerow(["بطيء الحركة",i[1],i[4],i[5]])
    elif rtype == "supplier_perf":
        writer.writerow(["المورد","أوامر الشراء","الإجمالي","الإستلامات","دقة التسليم","الجودة","المعدل"])
        sp_rows = db.session.query(
            Supplier.id, Supplier.name,
            func.count(PurchaseOrder.id).label("po_count"),
            func.sum(PurchaseOrder.total_amount).label("total_amount"),
            func.count(GoodsReceipt.id).label("grn_count"),
        ).outerjoin(PurchaseOrder, PurchaseOrder.supplier_id == Supplier.id
        ).outerjoin(GoodsReceipt, GoodsReceipt.po_id == PurchaseOrder.id
        ).group_by(Supplier.id).all()
        evals = SupplierEvaluation.query.with_entities(
            SupplierEvaluation.supplier_id,
            func.avg(SupplierEvaluation.quality).label("avg_quality"),
            func.avg(SupplierEvaluation.delivery).label("avg_delivery"),
        ).group_by(SupplierEvaluation.supplier_id).all()
        eval_map = {e[0]:{"quality":round(float(e[1] or 0),1),"delivery":round(float(e[2] or 0),1)} for e in evals}
        for r in sp_rows:
            ev = eval_map.get(r[0],{})
            avg = (ev.get("delivery",0) + ev.get("quality",0))/2
            writer.writerow([r[1] or "", r[2] or 0, round(float(r[3] or 0),2),
                           r[4] or 0, ev.get("delivery",0), ev.get("quality",0), round(avg,1)])
    elif rtype == "valuation":
        writer.writerow(["الصنف","الكود","الكمية الإجمالية","متوسط التكلفة","القيمة الإجمالية","عدد الطبقات"])
        for it in Item.query.filter_by(is_active=True).order_by(Item.name).all():
            v = InventoryLayer.get_valuation(it.id)
            if v["total_qty"] > 0:
                writer.writerow([it.name, it.code, v["total_qty"],
                               round(v["avg_cost"],2) if v["avg_cost"] else 0,
                               round(v["total_value"],2), v["layer_count"]])
    elif rtype == "audit":
        import csv as _csv
        filters = {}
        for k in ["user_id","action","resource","date_from","date_to","search"]:
            v = request.args.get(k)
            if v: filters[k] = v
        result = AuditService.get_logs(1, 99999, filters)
        writer.writerow(["#","التاريخ","المستخدم","الإجراء","العنصر","الوصف","IP","البيانات القديمة","البيانات الجديدة"])
        for l in (l.to_dict() for l in result["items"]):
            writer.writerow([l["id"],l["created_date"],l["user_name"],l["action"],
                            f"{l['resource']}#{l['resource_id']}" if l['resource_id'] else l['resource'],
                            l["description"] or "",l["ip_address"] or "",
                            l["old_data"] or "",l["new_data"] or ""])
    elif rtype == "abc":
        writer.writerow(["كود الصنف","الصنف","الفئة","الوحدة","إجمالي الكمية","سعر الوحدة","القيمة","%","تراكمي %","التصنيف"])
        with app.test_request_context():
            resp = abc_analysis()
            result = resp.get_json()["data"]
        for i in result.get("items", []):
            writer.writerow([i["item_code"], i["item_name"], i["category"], i["unit"],
                           i["total_stock"], i["unit_price"], i["total_value"],
                           f"{i['pct']}%", f"{i['cumulative_pct']}%", i["class"]])
    elif rtype == "aging":
        writer.writerow(["كود الصنف","الصنف","الوحدة","الكمية","القيمة","أقدم يوم","0-30","31-60","61-90","91-180","181-365","365+"])
        with app.test_request_context():
            resp = inventory_aging()
            result = resp.get_json()["data"]
        for i in result.get("items", []):
            b = i.get("qty_buckets",{})
            writer.writerow([i["item_code"], i["item_name"], i["unit"],
                           i["total_qty"], i["total_value"], i["oldest_days"],
                           b.get("0-30",0), b.get("31-60",0), b.get("61-90",0),
                           b.get("91-180",0), b.get("181-365",0), b.get("365+",0)])
    else:
        return err("نوع تقرير غير معروف")

    buf   = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    buf.seek(0)
    fname = f"warehouse_{rtype}_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return send_file(buf, mimetype="text/csv;charset=utf-8",
                     download_name=fname, as_attachment=True)


# ══════════════════════════════════════════════════════════════
#  PDF EXPORT  — GET /api/export/pdf?type=balance
# ══════════════════════════════════════════════════════════════
@api.route("/export/pdf", methods=["GET"])
@jwt_required()
def export_pdf():
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import (SimpleDocTemplate, Table,
                                         TableStyle, Paragraph, Spacer)
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    except ImportError:
        return err("مكتبة reportlab غير مثبتة — pip install reportlab")

    rtype  = request.args.get("type","balance")
    now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    buf = io.BytesIO()
    styles  = getSampleStyleSheet()
    t_style = ParagraphStyle("T", parent=styles["Title"],
                               fontSize=15, alignment=TA_CENTER, spaceAfter=10)
    f_style = ParagraphStyle("F", parent=styles["Normal"],
                               fontSize=8, alignment=TA_CENTER)
    elements = []
    col_w = None

    def make_table(data, headers, col_widths):
        tbl = Table([headers] + data, colWidths=col_widths, repeatRows=1)
        cmds = [
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1B4F72")),
            ("TEXTCOLOR", (0,0),(-1,0),colors.white),
            ("FONTSIZE",  (0,0),(-1,0),8),
            ("FONTSIZE",  (0,1),(-1,-1),7),
            ("ALIGN",     (0,0),(-1,-1),"CENTER"),
            ("VALIGN",    (0,0),(-1,-1),"MIDDLE"),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#EBF5FB")]),
            ("GRID",      (0,0),(-1,-1),0.4,colors.HexColor("#BDC3C7")),
            ("TOPPADDING",(0,0),(-1,-1),3),
            ("BOTTOMPADDING",(0,0),(-1,-1),3),
        ]
        tbl.setStyle(TableStyle(cmds))
        return tbl, cmds

    if rtype == "balance":
        whs   = Warehouse.query.filter_by(is_active=True).all()
        items = Item.query.filter_by(is_active=True).order_by(Item.name).all()
        headers = (["Code","Item Name","Category","Unit"] +
                   [w.name[:10] for w in whs] +
                   ["Total","Value SAR","Status"])
        data = []
        s_colors = {"critical":colors.HexColor("#FADBD8"),
                    "warning":colors.HexColor("#FDEBD0"),
                    "ok":colors.HexColor("#D5F5E3")}
        for item in items:
            stocks = {s.warehouse_id: s.quantity for s in item.stocks}
            total  = sum(stocks.values())
            st     = item.get_status()
            data.append(
                [item.code, item.name[:22],
                 item.category.name[:12] if item.category else "", item.unit] +
                [str(int(stocks.get(w.id,0))) for w in whs] +
                [str(int(total)), f"{total*item.unit_price:,.0f}",
                 {"critical":"Critical","warning":"Warning","ok":"OK"}[st]]
            )
        col_w = [2*cm, 4.5*cm, 2.5*cm, 1.5*cm] + [2*cm]*len(whs) + [2*cm, 2.8*cm, 2*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                 rightMargin=1.5*cm, leftMargin=1.5*cm,
                                 topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"Warehouse Balance Report — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        tbl, base_cmds = make_table(data, headers, col_w)
        status_col = len(headers) - 1
        for i, st in enumerate([it.get_status() for it in items], 1):
            tbl._argW[status_col] = 2*cm
            base_cmds.append(("BACKGROUND",(status_col,i),(status_col,i),s_colors[st]))
        tbl.setStyle(TableStyle(base_cmds))
        elements.append(tbl)
        total_val = sum(it.get_total_stock()*it.unit_price for it in items)
        elements.append(Spacer(1,0.4*cm))
        elements.append(Paragraph(f"Total Items: {len(items)} | Total Inventory Value: {total_val:,.0f} SAR",f_style))

    elif rtype == "movements":
        headers = ["#","Date","Time","Item","Type","Qty","Unit","Warehouse","User","Ref","Project","Engineer","Notes"]
        tmap = {"in":"In","out":"Out","return":"Return","damage":"Damage","transfer":"Transfer","adjustment":"Adj"}
        data = []
        for m in StockMovement.query.order_by(StockMovement.created_at.desc()).limit(500).all():
            data.append([
                str(m.id), m.created_at.strftime("%d/%m/%Y"), m.created_at.strftime("%H:%M"),
                m.item.name[:18] if m.item else "", tmap.get(m.type,m.type),
                str(m.quantity), m.item.unit if m.item else "",
                m.warehouse.name[:10] if m.warehouse else "",
                m.user.name[:10] if m.user else "",
                m.ref_number or "",
                (m.project or "")[:10],
                m.engineer_name or "",
                (m.notes or "")[:15],
            ])
        col_w = [1*cm,2*cm,1.5*cm,3.5*cm,1.5*cm,1.5*cm,1.2*cm,2*cm,2*cm,2*cm,2*cm,2*cm,2.5*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                 rightMargin=1*cm, leftMargin=1*cm,
                                 topMargin=1.5*cm, bottomMargin=1*cm)
        elements.append(Paragraph(f"Movements Report — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))
        elements.append(Spacer(1,0.3*cm))
        elements.append(Paragraph(f"Total Movements: {len(data)}", f_style))

    elif rtype == "transfers":
        headers = ["#","Ref#","Item","Qty","Unit","From","To","Requester","Date","Reason","Status"]
        data = []
        for t in Transfer.query.order_by(Transfer.created_at.desc()).all():
            data.append([
                str(t.id), t.ref_number or "",
                t.item.name[:18] if t.item else "", str(t.quantity),
                t.item.unit if t.item else "",
                t.from_wh.name[:10] if t.from_wh else "",
                t.to_wh.name[:10] if t.to_wh else "",
                t.requester.name[:10] if t.requester else "",
                t.created_at.strftime("%d/%m/%Y") if t.created_at else "",
                t.reason[:15] or "",
                t.STATUS_LABELS.get(t.status,t.status),
            ])
        col_w = [1*cm,2*cm,3.5*cm,1.5*cm,1.2*cm,2.5*cm,2.5*cm,2*cm,2*cm,3*cm,2*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                 rightMargin=1*cm, leftMargin=1*cm,
                                 topMargin=1.5*cm, bottomMargin=1*cm)
        elements.append(Paragraph(f"Transfers Report — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))
        elements.append(Spacer(1,0.3*cm))
        elements.append(Paragraph(f"Total Transfers: {len(data)}", f_style))

    elif rtype == "diffs":
        headers = ["#","Item","Code","Warehouse","System","Actual","Diff","Date"]
        data = []
        diffs = (db.session.query(InventoryCountLine)
                 .join(InventoryCount)
                 .filter(
                     InventoryCountLine.actual_quantity.isnot(None),
                     InventoryCountLine.actual_quantity != InventoryCountLine.system_quantity,
                     InventoryCount.status == "completed"
                 )
                 .order_by(InventoryCount.completed_at.desc())
                 .all())
        for l in diffs:
            data.append([
                str(l.id),
                l.item.name[:18] if l.item else "",
                l.item.code[:10] if l.item else "",
                l.count.warehouse.name[:10] if l.count and l.count.warehouse else "",
                str(l.system_quantity), str(l.actual_quantity),
                str(l.difference),
                l.count.completed_at.strftime("%d/%m/%Y") if l.count and l.count.completed_at else "",
            ])
        col_w = [1*cm,3.5*cm,2*cm,2.5*cm,2*cm,2*cm,2*cm,2*cm]
        doc = SimpleDocTemplate(buf, pagesize=A4,
                                 rightMargin=1.5*cm, leftMargin=1.5*cm,
                                 topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"Count Differences Report — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))
        elements.append(Spacer(1,0.3*cm))
        elements.append(Paragraph(f"Total Differences: {len(data)}", f_style))

    elif rtype == "budget":
        headers = ["Project","Code","Client","Budget","Spent","Remaining","Usage %","Status"]
        data = []
        bd = ReportService.budget()
        for p in bd["projects"]:
            data.append([p["name"][:18],p["code"],p["client"][:12],
                        f"{p['budget']:,.0f}",f"{p['spent']:,.0f}",
                        f"{p['remaining']:,.0f}",f"{p['usage_pct']}%",
                        p["status_label"]])
        col_w = [3.5*cm,2*cm,2.5*cm,2.5*cm,2.5*cm,2.5*cm,2*cm,1.5*cm]
        doc = SimpleDocTemplate(buf, pagesize=A4,
                                 rightMargin=1.5*cm, leftMargin=1.5*cm,
                                 topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"Projects Budget Report — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))
        s = bd["summary"]
        elements.append(Spacer(1,0.3*cm))
        elements.append(Paragraph(
            f"Total Projects: {s['total']} | Active: {s['active']} | "
            f"Total Budget: {s['total_budget']:,.0f} SAR | "
            f"Total Spent: {s['total_spent']:,.0f} SAR | "
            f"Remaining: {s['total_remaining']:,.0f} SAR", f_style))

    elif rtype == "consumption":
        days = int(request.args.get("days",30))
        since = datetime.datetime.utcnow() - datetime.timedelta(days=days)
        limit = int(request.args.get("limit",50))
        headers = ["Item","Code","Qty Consumed","Value","Movements","Avg/Day"]
        data = []
        rows = db.session.query(
            StockMovement.item_id, Item.name, Item.code,
            func.sum(StockMovement.quantity).label("total_qty"),
            func.sum(StockMovement.quantity * StockMovement.unit_price).label("total_val"),
            func.count(StockMovement.id).label("mov_count")
        ).join(Item, StockMovement.item_id == Item.id
        ).filter(StockMovement.type.in_(["out","transfer","damage"]),
                 StockMovement.created_at >= since
        ).group_by(StockMovement.item_id
        ).order_by(func.sum(StockMovement.quantity).desc()
        ).limit(limit).all()
        for r in rows:
            data.append([r[1][:22], r[2], str(int(r[3])),
                        f"{round(float(r[4] or 0),2):,.2f}",
                        str(r[5]), f"{round(float(r[3])/max(days,1),2):,.2f}"])
        col_w = [4.5*cm, 2.5*cm, 2.5*cm, 3*cm, 2.5*cm, 2.5*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                 rightMargin=1.5*cm, leftMargin=1.5*cm,
                                 topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"Consumption Report ({days}d) — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))
        elements.append(Spacer(1,0.3*cm))
        total_q = sum(float(r[3]) for r in rows) if rows else 0
        elements.append(Paragraph(f"Total Items: {len(data)} | Total Consumed: {int(total_q)} units", f_style))

    elif rtype == "fastslow":
        days = int(request.args.get("days",90))
        since = datetime.datetime.utcnow() - datetime.timedelta(days=days)
        limit = int(request.args.get("limit",50))
        headers = ["Type","Item","Consumed","Current Stock"]
        data = []
        out_q = db.session.query(StockMovement.item_id, func.sum(StockMovement.quantity).label("qty")).filter(
            StockMovement.type.in_(["out","transfer","damage"]), StockMovement.created_at >= since
        ).group_by(StockMovement.item_id).subquery()
        items = db.session.query(Item.id, Item.name, Item.code,
                                 func.coalesce(out_q.c.qty, 0).label("consumed"),
                                 func.coalesce(func.sum(Stock.quantity), 0).label("current_stock")
        ).outerjoin(out_q, Item.id == out_q.c.item_id
        ).outerjoin(Stock, Stock.item_id == Item.id
        ).group_by(Item.id).all()
        fast = sorted(items, key=lambda x: float(x.consumed or 0), reverse=True)[:limit]
        slow = sorted([it for it in items if float(it.consumed or 0) <= 0], key=lambda x: float(x.current_stock or 0), reverse=True)[:limit]
        for i in fast:
            data.append(["Fast", i[2], str(int(i[3])), str(int(i[4]))])
        data.append(["—","—","—","—"])
        for i in slow:
            data.append(["Slow", i[2], str(int(i[3])), str(int(i[4]))])
        col_w = [2*cm, 4.5*cm, 3*cm, 3*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                 rightMargin=1.5*cm, leftMargin=1.5*cm,
                                 topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"Fast/Slow Moving Items ({days}d) — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))

    elif rtype == "supplier_perf":
        headers = ["Supplier","PO Count","Total Amount","GRN Count","Delivery","Quality","Avg Rating"]
        data = []
        sp_rows = db.session.query(
            Supplier.id, Supplier.name,
            func.count(PurchaseOrder.id).label("po_count"),
            func.sum(PurchaseOrder.total_amount).label("total_amount"),
            func.count(GoodsReceipt.id).label("grn_count"),
        ).outerjoin(PurchaseOrder, PurchaseOrder.supplier_id == Supplier.id
        ).outerjoin(GoodsReceipt, GoodsReceipt.po_id == PurchaseOrder.id
        ).group_by(Supplier.id).all()
        sp_evals = SupplierEvaluation.query.with_entities(
            SupplierEvaluation.supplier_id,
            func.avg(SupplierEvaluation.quality).label("avg_quality"),
            func.avg(SupplierEvaluation.delivery).label("avg_delivery"),
        ).group_by(SupplierEvaluation.supplier_id).all()
        sp_emap = {e[0]:{"quality":round(float(e[1] or 0),1),"delivery":round(float(e[2] or 0),1)} for e in sp_evals}
        for r in sp_rows:
            ev = sp_emap.get(r[0],{})
            avg = (ev.get("delivery",0) + ev.get("quality",0))/2
            data.append([(r[1] or "")[:20], str(r[2] or 0),
                        f"{round(float(r[3] or 0),2):,.2f}", str(r[4] or 0),
                        str(ev.get("delivery",0)), str(ev.get("quality",0)),
                        f"{round(avg,1)}"])
        col_w = [4*cm, 2.5*cm, 3*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2.5*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                 rightMargin=1.5*cm, leftMargin=1.5*cm,
                                 topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"Supplier Performance Report — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))

    elif rtype == "valuation":
        headers = ["Item","Code","Total Qty","Avg Cost","Total Value","Layers"]
        data = []
        for it in Item.query.filter_by(is_active=True).order_by(Item.name).all():
            v = InventoryLayer.get_valuation(it.id)
            if v["total_qty"] > 0:
                data.append([it.name[:22], it.code, str(v["total_qty"]),
                           f"{round(v['avg_cost'],2):,.2f}" if v['avg_cost'] else "0",
                           f"{round(v['total_value'],2):,.2f}", str(v["layer_count"])])
        col_w = [4.5*cm, 2.5*cm, 2.5*cm, 2.5*cm, 3*cm, 2*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                 rightMargin=1.5*cm, leftMargin=1.5*cm,
                                 topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"FIFO Inventory Valuation — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))
        grand_v = sum(v["total_value"] for v in [InventoryLayer.get_valuation(it.id) for it in Item.query.filter_by(is_active=True).all()] if v["total_qty"] > 0)
        elements.append(Spacer(1,0.3*cm))
        elements.append(Paragraph(f"Total Items: {len(data)} | Grand Total Value: {grand_v:,.2f} SAR", f_style))

    elif rtype == "audit":
        headers = ["#","Date","User","Action","Resource","Description","IP"]
        data = []
        filters = {}
        for k in ["user_id","action","resource","date_from","date_to","search"]:
            v = request.args.get(k)
            if v: filters[k] = v
        result = AuditService.get_logs(1, 5000, filters)
        for l in result["items"]:
            data.append([
                str(l.id), (l.created_date or "")[:16],
                (l.user_name or "")[:15], l.action,
                f"{l.resource}#{l.resource_id}" if l.resource_id else (l.resource or ""),
                (l.description or "")[:25], l.ip_address or "",
            ])
        col_w = [1*cm, 2.5*cm, 2.5*cm, 1.8*cm, 3*cm, 5*cm, 2.5*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                rightMargin=1.5*cm, leftMargin=1.5*cm,
                                topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"Audit Log Report — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))
        elements.append(Spacer(1,0.3*cm))
        elements.append(Paragraph(f"Total Logs: {len(data)}", f_style))

    elif rtype == "abc":
        headers = ["Code","Item","Category","Unit","Stock","Unit Price","Value","%","Cum %","Class"]
        data = []
        with app.test_request_context():
            resp = abc_analysis()
            result = resp.get_json()["data"]
        for i in result.get("items", []):
            data.append([
                i["item_code"], i["item_name"][:20], i["category"][:12], i["unit"],
                str(int(i["total_stock"])), f"{i['unit_price']:,.2f}",
                f"{i['total_value']:,.2f}", f"{i['pct']}%",
                f"{i['cumulative_pct']}%", i["class"],
            ])
        col_w = [2*cm, 4*cm, 2.5*cm, 1.5*cm, 1.5*cm, 2.5*cm, 3*cm, 1.5*cm, 1.5*cm, 1.5*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                rightMargin=1.5*cm, leftMargin=1.5*cm,
                                topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"ABC Inventory Analysis — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))
        s = result.get("summary", {})
        elements.append(Spacer(1,0.3*cm))
        elements.append(Paragraph(
            f"A: {s.get('A',{}).get('count',0)} items ({s.get('A',{}).get('value',0):,.0f} SAR) | "
            f"B: {s.get('B',{}).get('count',0)} items ({s.get('B',{}).get('value',0):,.0f} SAR) | "
            f"C: {s.get('C',{}).get('count',0)} items ({s.get('C',{}).get('value',0):,.0f} SAR) | "
            f"Total: {s.get('grand_total',0):,.0f} SAR", f_style))

    elif rtype == "aging":
        headers = ["Code","Item","Unit","Qty","Value","Oldest Day","0-30","31-60","61-90","91-180","181-365","365+"]
        data = []
        with app.test_request_context():
            resp = inventory_aging()
            result = resp.get_json()["data"]
        for i in result.get("items", []):
            b = i.get("qty_buckets",{})
            data.append([
                i["item_code"], i["item_name"][:20], i["unit"],
                str(int(i["total_qty"])), f"{i['total_value']:,.2f}", f"{i['oldest_days']}d",
                str(int(b.get("0-30",0))), str(int(b.get("31-60",0))),
                str(int(b.get("61-90",0))), str(int(b.get("91-180",0))),
                str(int(b.get("181-365",0))), str(int(b.get("365+",0))),
            ])
        col_w = [2*cm, 3.5*cm, 1.5*cm, 1.5*cm, 2.5*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm]
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                                rightMargin=1.5*cm, leftMargin=1.5*cm,
                                topMargin=2*cm, bottomMargin=1.5*cm)
        elements.append(Paragraph(f"Inventory Aging Report — {now_str}", t_style))
        elements.append(Spacer(1, 0.3*cm))
        elements.append(make_table(data, headers, col_w))

    else:
        return err("نوع تقرير غير معروف")

    doc.build(elements)
    buf.seek(0)
    fname = f"warehouse_{rtype}_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
    return send_file(buf, mimetype="application/pdf",
                     download_name=fname, as_attachment=True)

# ══════════════════════════════════════════════════════════════
#  FILE UPLOAD
# ══════════════════════════════════════════════════════════════
@api.route("/upload", methods=["POST"])
@jwt_required()
def upload_file():
    uid = int(get_jwt_identity())
    item_id = request.form.get("item_id", type=int)
    if not item_id: return err("معرف الصنف مطلوب")
    item = Item.query.get_or_404(item_id)
    f = request.files.get("file")
    if not f: return err("الملف مطلوب")
    ext = f.filename.rsplit(".",1)[-1].lower() if "." in f.filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        return err("نوع الملف غير مسموح: يرجى رفع PDF, صور, مستندات Office أو CSV/ZIP فقط")
    f.seek(0, os.SEEK_END)
    fsize = f.tell()
    f.seek(0)
    if fsize > MAX_FILE_SIZE:
        return err("حجم الملف يتجاوز 16 ميجابايت")
    ts  = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = f"item_{item_id}_{ts}_{secure_filename(f.filename)}"
    folder = os.path.join(current_app.root_path, "..", "uploads")
    os.makedirs(folder, exist_ok=True)
    f.save(os.path.join(folder, safe_name))
    mime_type = mimetypes.guess_type(f.filename)[0] or "application/octet-stream"
    attach = ItemAttachment(
        item_id=item_id, filename=safe_name,
        original_name=f.filename, file_type=mime_type,
        file_size=os.path.getsize(os.path.join(folder, safe_name)),
        uploaded_by=uid)
    db.session.add(attach)
    AuditService.log("upload","item_attachment",attach.id,f"رفع ملف: {f.filename} للصنف {item.name}")
    db.session.commit()
    return created(attach.to_dict(), "تم رفع الملف")

@api.route("/items/<int:iid>/attachments", methods=["GET"])
@jwt_required()
def get_item_attachments(iid):
    files = ItemAttachment.query.filter_by(item_id=iid).order_by(ItemAttachment.created_at.desc()).all()
    return ok([f.to_dict() for f in files])

@api.route("/uploads/<path:filename>", methods=["GET"])
@jwt_required()
def serve_upload(filename):
    folder = os.path.join(current_app.root_path, "..", "uploads")
    return send_from_directory(folder, filename)

# ══════════════════════════════════════════════════════════════
#  NOTIFICATION SETTINGS
# ══════════════════════════════════════════════════════════════
@api.route("/settings/notifications", methods=["GET"])
@jwt_required()
def get_notif_settings():
    configs = NotificationConfig.query.all()
    d = {c.key: c.value for c in configs}
    return ok({
        "whatsapp_enabled": d.get("whatsapp_enabled","false"),
        "whatsapp_number": d.get("whatsapp_number",""),
        "whatsapp_cloud_api": d.get("whatsapp_cloud_api","false"),
        "whatsapp_token": d.get("whatsapp_token",""),
        "whatsapp_phone_id": d.get("whatsapp_phone_id",""),
        "email_enabled": d.get("email_enabled","false"),
        "smtp_host": d.get("smtp_host",""),
        "smtp_port": d.get("smtp_port","587"),
        "smtp_user": d.get("smtp_user",""),
        "smtp_pass": d.get("smtp_pass",""),
        "email_from": d.get("email_from",""),
        "notify_low_stock": d.get("notify_low_stock","true"),
        "notify_transfer": d.get("notify_transfer","true"),
    })

@api.route("/settings/notifications", methods=["POST"])
@require_role("admin","manager")
def save_notif_settings():
    data = request.get_json() or {}
    for key in ["whatsapp_enabled","whatsapp_number","whatsapp_cloud_api","whatsapp_token","whatsapp_phone_id",
                "email_enabled","smtp_host","smtp_port","smtp_user","smtp_pass",
                "email_from","notify_low_stock","notify_transfer"]:
        if key in data:
            c = NotificationConfig.query.filter_by(key=key).first()
            if not c:
                c = NotificationConfig(key=key, value=str(data[key]))
                db.session.add(c)
            else:
                c.value = str(data[key])
    db.session.commit()
    return ok(message="تم حفظ إعدادات الإشعارات")

@api.route("/notifications/test-email", methods=["POST"])
@require_role("admin","manager")
def test_email():
    data = request.get_json() or {}
    to = data.get("to","")
    if not to: return err("البريد المستلم مطلوب")
    try:
        host = NotificationConfig.query.filter_by(key="smtp_host").first()
        port = NotificationConfig.query.filter_by(key="smtp_port").first()
        user = NotificationConfig.query.filter_by(key="smtp_user").first()
        pwd  = NotificationConfig.query.filter_by(key="smtp_pass").first()
        frm  = NotificationConfig.query.filter_by(key="email_from").first()
        if not host or not host.value: return err("SMTP غير مهيأ")
        msg = MIMEMultipart()
        msg["From"] = frm.value if frm else user.value if user else "noreply@wms.local"
        msg["To"] = to
        msg["Subject"] = "🧪 اختبار إعدادات البريد - نظام إدارة المخازن"
        msg.attach(MIMEText("تم إعداد البريد الإلكتروني بنجاح ✅\n\nنظام إدارة المخازن", "plain", "utf-8"))
        with smtplib.SMTP(host.value, int(port.value if port else 587), timeout=10) as s:
            s.starttls()
            if user and user.value and pwd and pwd.value:
                s.login(user.value, pwd.value)
            s.send_message(msg)
        return ok(message="✅ تم إرسال بريد الاختبار بنجاح")
    except Exception as e:
        return err(f"❌ فشل الإرسال: {str(e)[:100]}")

@api.route("/notifications/test-whatsapp", methods=["POST"])
@require_role("admin","manager")
def test_whatsapp():
    data = request.get_json() or {}
    number = data.get("number","")
    if not number: return err("رقم الواتساب مطلوب")
    number = number.strip().replace(" ","").replace("-","")
    if not number.startswith("+"): number = "+" + number
    from app.services.whatsapp import WhatsAppService
    ok, result = WhatsAppService.send_notification(number, "🔔 اختبار إشعار", "هذه رسالة اختبارية من نظام النبلاء")
    if result.get("method") == "cloud_api":
        return ok(result, message="✅ تم إرسال رسالة واتساب عبر Cloud API")
    return ok(result, message="✅ رابط واتساب جاهز (الوضع التجريبي)")

# ══════════════════════════════════════════════════════════════
#  BACKUP
# ══════════════════════════════════════════════════════════════
@api.route("/settings/backup", methods=["GET"])
@jwt_required()
def get_backup_settings():
    configs = BackupConfig.query.all()
    d = {c.key: c.value for c in configs}
    return ok({
        "auto_backup_enabled": d.get("auto_backup_enabled","false"),
        "backup_interval_hours": d.get("backup_interval_hours","24"),
        "backup_keep_count": d.get("backup_keep_count","10"),
        "last_backup": d.get("last_backup",""),
    })

@api.route("/settings/backup", methods=["POST"])
@require_role("admin")
def save_backup_settings():
    data = request.get_json() or {}
    for key in ["auto_backup_enabled","backup_interval_hours","backup_keep_count","last_backup"]:
        if key in data:
            c = BackupConfig.query.filter_by(key=key).first()
            if not c:
                c = BackupConfig(key=key, value=str(data[key]))
                db.session.add(c)
            else:
                c.value = str(data[key])
    db.session.commit()
    return ok(message="تم حفظ إعدادات النسخ الاحتياطي")

@api.route("/backup/trigger", methods=["POST"])
@require_role("admin")
def trigger_backup():
    try:
        db_path = current_app.config.get("SQLALCHEMY_DATABASE_URI","").replace("sqlite:///","")
        if not os.path.isabs(db_path):
            db_path = os.path.join(current_app.root_path, "..", db_path)
        if not os.path.exists(db_path): return err("قاعدة البيانات غير موجودة")
        backup_dir = os.path.join(current_app.root_path, "..", "backups")
        os.makedirs(backup_dir, exist_ok=True)
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_name = f"wms_backup_{ts}.db"
        import shutil
        shutil.copy2(db_path, os.path.join(backup_dir, backup_name))
        keep = int((BackupConfig.query.filter_by(key="backup_keep_count").first() or BackupConfig(key="backup_keep_count",value="10")).value)
        all_backups = sorted([f for f in os.listdir(backup_dir) if f.startswith("wms_backup_")], reverse=True)
        for old in all_backups[keep:]:
            try: os.remove(os.path.join(backup_dir, old))
            except: pass
        c = BackupConfig.query.filter_by(key="last_backup").first()
        if not c:
            c = BackupConfig(key="last_backup", value=ts)
            db.session.add(c)
        else:
            c.value = ts
        AuditService.log("backup","system",desc=f"نسخة احتياطية: {backup_name}")
        db.session.commit()
        return ok({"filename":backup_name}, message=f"✅ تم إنشاء النسخة الاحتياطية {backup_name}")
    except Exception as e:
        return err(f"❌ فشل النسخ الاحتياطي: {str(e)[:200]}")

@api.route("/backups", methods=["GET"])
@require_role("admin")
def list_backups():
    backup_dir = os.path.join(os.path.dirname(current_app.root_path), "backups")
    os.makedirs(backup_dir, exist_ok=True)
    files = []
    for f in sorted(os.listdir(backup_dir), reverse=True):
        fp = os.path.join(backup_dir, f)
        if os.path.isfile(fp):
            sz = os.path.getsize(fp)
            files.append({"name":f,"size":sz,
                "size_label":f"{sz/1024/1024:.1f} MB" if sz>1024*1024 else f"{sz/1024:.0f} KB",
                "created":datetime.datetime.fromtimestamp(os.path.getmtime(fp)).isoformat()})
    return ok(files)

@api.route("/backups/<path:name>", methods=["GET"])
@require_role("admin")
def download_backup(name):
    backup_dir = os.path.join(os.path.dirname(current_app.root_path), "backups")
    return send_from_directory(backup_dir, name, as_attachment=True)

@api.route("/settings/auto-reorder", methods=["GET"])
@require_role("admin","manager")
def get_auto_reorder_settings():
    cfg = {}
    for key in ["auto_reorder_enabled","auto_reorder_supplier_id","auto_reorder_warehouse_id"]:
        c = BackupConfig.query.filter_by(key=key).first()
        cfg[key] = c.value if c else ""
    suppliers = [{"id":s.id,"name":s.name} for s in Supplier.query.filter_by(is_active=True).all()]
    warehouses = [{"id":w.id,"name":w.name} for w in Warehouse.query.filter_by(is_active=True).all()]
    return ok({"config": cfg, "suppliers": suppliers, "warehouses": warehouses})

@api.route("/settings/auto-reorder", methods=["POST"])
@require_role("admin","manager")
def update_auto_reorder_settings():
    data = request.get_json() or {}
    for key in ["auto_reorder_enabled","auto_reorder_supplier_id","auto_reorder_warehouse_id"]:
        if key in data:
            c = BackupConfig.query.filter_by(key=key).first()
            if not c:
                c = BackupConfig(key=key, value=str(data[key]))
                db.session.add(c)
            else:
                c.value = str(data[key])
    db.session.commit()
    return ok(message="✅ تم حفظ إعدادات التوريد التلقائي")

@api.route("/auto-reorder/run", methods=["POST"])
@require_role("admin","manager")
def run_auto_reorder():
    from app import auto_reorder_check
    result = auto_reorder_check()
    return ok(result, message=f"✅ تم إنشاء {result.get('created',0)} طلب توريد")

@api.route("/backup/restore", methods=["POST"])
@require_role("admin")
def restore_backup():
    """
    Restore database from backup
    ---
    parameters:
      - name: body
        in: body
        required: true
        schema:
          properties:
            name: {type: string, example: wms_backup_20250101_120000.db}
    responses:
      200: {description: Database restored}
    """
    data = request.get_json() or {}
    name = data.get("name","")
    if not name: return err("يجب تحديد اسم ملف الاستعادة")
    backup_dir = os.path.join(os.path.dirname(current_app.root_path), "backups")
    backup_path = os.path.join(backup_dir, name)
    if not os.path.isfile(backup_path): return err("ملف الاستعادة غير موجود")
    db_path = current_app.config.get("SQLALCHEMY_DATABASE_URI","").replace("sqlite:///","")
    if not db_path: return err("غير مدعوم لقواعد البيانات غير SQLite")
    db_path = os.path.join(os.path.dirname(current_app.root_path), db_path)
    try:
        import shutil
        shutil.copy2(backup_path, db_path)
        AuditService.log("restore","database",0,f"استعادة قاعدة البيانات من: {name}")
        return ok(message=f"تمت استعادة قاعدة البيانات من {name}")
    except Exception as e:
        return err(f"فشلت الاستعادة: {str(e)}")

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — PURCHASE REQUESTS
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/pr", methods=["GET"])
@jwt_required()
def get_pr_list():
    search = request.args.get("search","")
    status = request.args.get("status","")
    q = PurchaseRequest.query
    if search: q = q.filter(PurchaseRequest.ref_number.contains(search) | PurchaseRequest.department.contains(search))
    if status: q = q.filter(PurchaseRequest.status == status)
    return ok([p.to_dict() for p in q.order_by(PurchaseRequest.created_at.desc()).all()])

@api.route("/procurement/pr", methods=["POST"])
@require_role("admin","manager","keeper")
def create_pr():
    data = request.get_json() or {}
    errs = validate(data, ["requester_name"])
    if errs: return err("تحقق من الحقول المطلوبة", errors=errs)
    ref = gen_ref("PR")
    pr = PurchaseRequest(ref_number=ref, requester_id=g.current_user.id,
        department=data.get("department",""), priority=data.get("priority","medium"),
        notes=data.get("notes",""), status="draft")
    if "project_id" in data: pr.project_id = int(data["project_id"])
    if "required_date" in data: pr.required_date = parse_date(data["required_date"])
    if "request_date" in data: pr.request_date = parse_date(data["request_date"])
    db.session.add(pr); db.session.flush()
    for it in (data.get("items") or []):
        pi = PRItem(pr_id=pr.id, item_name=it.get("item_name",""),
            category=it.get("category",""), unit=it.get("unit",""),
            quantity=float(it.get("quantity",0)), current_stock=float(it.get("current_stock",0)),
            min_stock=float(it.get("min_stock",0)), estimated_cost=float(it.get("estimated_cost",0)),
            notes=it.get("notes",""))
        db.session.add(pi)
    AuditService.log("create","purchase_request",pr.id,f"إنشاء طلب شراء: {ref}")
    db.session.commit()
    return created(pr.to_dict())

@api.route("/procurement/pr/<int:pid>", methods=["GET"])
@jwt_required()
def get_pr(pid):
    pr = PurchaseRequest.query.get_or_404(pid)
    return ok(pr.to_dict())

@api.route("/procurement/pr/<int:pid>", methods=["PUT"])
@require_role("admin","manager")
def update_pr(pid):
    pr = PurchaseRequest.query.get_or_404(pid)
    data = request.get_json() or {}
    for f in ("department","priority","notes"):
        if f in data: setattr(pr, f, data[f])
    if "project_id" in data: pr.project_id = int(data["project_id"]) if data["project_id"] else None
    if "required_date" in data: pr.required_date = parse_date(data["required_date"])
    if "items" in data:
        PRItem.query.filter_by(pr_id=pid).delete()
        for it in data["items"]:
            pi = PRItem(pr_id=pid, item_name=it.get("item_name",""),
                category=it.get("category",""), unit=it.get("unit",""),
                quantity=float(it.get("quantity",0)), current_stock=float(it.get("current_stock",0)),
                min_stock=float(it.get("min_stock",0)), estimated_cost=float(it.get("estimated_cost",0)),
                notes=it.get("notes",""))
            db.session.add(pi)
    AuditService.log("edit","purchase_request",pid,f"تعديل طلب شراء: {pr.ref_number}")
    db.session.commit()
    return ok(pr.to_dict())

@api.route("/procurement/pr/<int:pid>/submit", methods=["POST"])
@require_role("admin","manager","keeper")
def submit_pr(pid):
    pr = PurchaseRequest.query.get_or_404(pid)
    if pr.status != "draft": return err("يمكن تقديم المسودات فقط")
    if not pr.items: return err("يجب إضافة أصناف على الأقل")
    pr.status = "pending"
    # notification to managers
    admins = User.query.filter(User.role.in_(["admin","manager"])).all()
    for u in admins:
        db.session.add(Notification(type="approval", title="طلب شراء جديد",
            message=f"طلب شراء {pr.ref_number} بانتظار الموافقة",
            user_id=u.id, ref_type="pr", ref_id=pr.id))
    AuditService.log("submit","purchase_request",pid,f"تقديم طلب شراء: {pr.ref_number}")
    db.session.commit()
    from app import get_socketio
    sio = get_socketio()
    for u in admins:
        sio.emit("notification_update", {"unread": 1}, to=f"user_{u.id}")
    return ok(pr.to_dict())

@api.route("/procurement/pr/<int:pid>/approve", methods=["POST"])
@require_role("admin","manager")
def approve_pr(pid):
    pr = PurchaseRequest.query.get_or_404(pid)
    if pr.status != "pending": return err("يمكن اعتماد الطلبات قيد الانتظار فقط")
    data = request.get_json() or {}
    level = int(data.get("level",1))
    pr.status = "approved"
    al = ApprovalLog(resource_type="pr", resource_id=pid, level=level,
        approver_id=g.current_user.id, status="approved", notes=data.get("notes",""))
    db.session.add(al)
    if pr.requester_id:
        db.session.add(Notification(type="pr_approved",
            title="تم اعتماد طلب الشراء", message=f"تم اعتماد {pr.ref_number}",
            user_id=pr.requester_id, ref_type="pr", ref_id=pr.id))
    AuditService.log("approve","purchase_request",pid,f"اعتماد طلب شراء: {pr.ref_number}")
    db.session.commit()
    from app import get_socketio
    sio = get_socketio()
    if pr.requester_id:
        sio.emit("notification_update", {"unread": 1}, to=f"user_{pr.requester_id}")
    return ok(pr.to_dict())

@api.route("/procurement/pr/<int:pid>/reject", methods=["POST"])
@require_role("admin","manager")
def reject_pr(pid):
    pr = PurchaseRequest.query.get_or_404(pid)
    if pr.status != "pending": return err("يمكن رفض الطلبات قيد الانتظار فقط")
    data = request.get_json() or {}
    level = int(data.get("level",1))
    pr.status = "rejected"
    al = ApprovalLog(resource_type="pr", resource_id=pid, level=level,
        approver_id=g.current_user.id, status="rejected", notes=data.get("notes",""))
    db.session.add(al)
    AuditService.log("reject","purchase_request",pid,f"رفض طلب شراء: {pr.ref_number}")
    db.session.commit()
    return ok(pr.to_dict())

@api.route("/procurement/pr/<int:pid>", methods=["DELETE"])
@require_role("admin")
def delete_pr(pid):
    pr = PurchaseRequest.query.get_or_404(pid)
    db.session.delete(pr)
    AuditService.log("delete","purchase_request",pid,f"حذف طلب شراء: {pr.ref_number}")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — APPROVAL LOG
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/approvals/<string:rtype>/<int:rid>", methods=["GET"])
@jwt_required()
def get_approvals(rtype, rid):
    logs = ApprovalLog.query.filter_by(resource_type=rtype, resource_id=rid).order_by(ApprovalLog.created_at.desc()).all()
    return ok([l.to_dict() for l in logs])

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — RFQ
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/rfq", methods=["GET"])
@jwt_required()
def get_rfq_list():
    search = request.args.get("search","")
    status = request.args.get("status","")
    q = RFQ.query
    if search: q = q.filter(RFQ.ref_number.contains(search))
    if status: q = q.filter(RFQ.status == status)
    return ok([r.to_dict() for r in q.order_by(RFQ.created_at.desc()).all()])

@api.route("/procurement/rfq", methods=["POST"])
@require_role("admin","manager")
def create_rfq():
    data = request.get_json() or {}
    ref = gen_ref("RFQ")
    rfq = RFQ(ref_number=ref, pr_id=data.get("pr_id"),
        delivery_terms=data.get("delivery_terms",""), payment_terms=data.get("payment_terms",""),
        warranty=data.get("warranty",""), notes=data.get("notes",""), status="draft")
    if "valid_until" in data: rfq.valid_until = parse_date(data["valid_until"])
    db.session.add(rfq); db.session.flush()
    for sid in (data.get("supplier_ids") or []):
        db.session.add(RFQSupplier(rfq_id=rfq.id, supplier_id=int(sid)))
    for it in (data.get("items") or []):
        db.session.add(RFQItem(rfq_id=rfq.id, item_name=it.get("item_name",""),
            quantity=float(it.get("quantity",0)), unit=it.get("unit",""), notes=it.get("notes","")))
    AuditService.log("create","rfq",rfq.id,f"إنشاء طلب عرض سعر: {ref}")
    db.session.commit()
    return created(rfq.to_dict())

@api.route("/procurement/rfq/<int:rid>", methods=["GET"])
@jwt_required()
def get_rfq(rid):
    rfq = RFQ.query.get_or_404(rid)
    return ok(rfq.to_dict())

@api.route("/procurement/rfq/<int:rid>", methods=["PUT"])
@require_role("admin","manager")
def update_rfq(rid):
    rfq = RFQ.query.get_or_404(rid)
    data = request.get_json() or {}
    for f in ("delivery_terms","payment_terms","warranty","notes","status"):
        if f in data: setattr(rfq, f, data[f])
    if "valid_until" in data: rfq.valid_until = parse_date(data["valid_until"])
    if "supplier_ids" in data:
        RFQSupplier.query.filter_by(rfq_id=rid).delete(); db.session.flush()
        for sid in data["supplier_ids"]:
            db.session.add(RFQSupplier(rfq_id=rid, supplier_id=int(sid)))
    AuditService.log("edit","rfq",rid,f"تعديل طلب عرض سعر: {rfq.ref_number}")
    db.session.commit()
    return ok(rfq.to_dict())

@api.route("/procurement/rfq/<int:rid>/send", methods=["POST"])
@require_role("admin","manager")
def send_rfq(rid):
    rfq = RFQ.query.get_or_404(rid)
    rfq.status = "sent"
    AuditService.log("send","rfq",rid,f"إرسال طلب عرض سعر: {rfq.ref_number}")
    db.session.commit()
    return ok(rfq.to_dict())

@api.route("/procurement/rfq/<int:rid>", methods=["DELETE"])
@require_role("admin")
def delete_rfq(rid):
    rfq = RFQ.query.get_or_404(rid)
    db.session.delete(rfq)
    AuditService.log("delete","rfq",rid,f"حذف طلب عرض سعر")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — QUOTATIONS
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/quotations", methods=["GET"])
@jwt_required()
def get_quotations():
    rfq_id = request.args.get("rfq_id")
    q = Quotation.query
    if rfq_id: q = q.filter(Quotation.rfq_id == int(rfq_id))
    return ok([qt.to_dict() for qt in q.order_by(Quotation.created_at.desc()).all()])

@api.route("/procurement/quotations", methods=["POST"])
@require_role("admin","manager")
def create_quotation():
    data = request.get_json() or {}
    ref = gen_ref("QT")
    qt = Quotation(ref_number=ref, rfq_id=data.get("rfq_id"), supplier_id=data.get("supplier_id"),
        amount=float(data.get("amount",0)), vat=float(data.get("vat",0)),
        delivery_cost=float(data.get("delivery_cost",0)),
        delivery_days=int(data.get("delivery_days",0)), warranty_period=data.get("warranty_period",""),
        notes=data.get("notes",""), status="pending")
    if "valid_until" in data: qt.valid_until = parse_date(data["valid_until"])
    qt.total = qt.amount + qt.vat + qt.delivery_cost
    db.session.add(qt); db.session.flush()
    for it in (data.get("items") or []):
        qty = float(it.get("quantity",0)); up = float(it.get("unit_price",0))
        db.session.add(QuotationItem(quotation_id=qt.id, item_name=it.get("item_name",""),
            quantity=qty, unit_price=up, total=round(qty*up,2)))
    AuditService.log("create","quotation",qt.id,f"إنشاء عرض سعر: {ref}")
    db.session.commit()
    return created(qt.to_dict())

@api.route("/procurement/quotations/<int:qid>", methods=["GET"])
@jwt_required()
def get_quotation(qid):
    qt = Quotation.query.get_or_404(qid)
    return ok(qt.to_dict())

@api.route("/procurement/quotations/<int:qid>", methods=["PUT"])
@require_role("admin","manager")
def update_quotation(qid):
    qt = Quotation.query.get_or_404(qid)
    data = request.get_json() or {}
    for f in ("amount","vat","delivery_cost","delivery_days","warranty_period","notes","status"):
        if f in data: setattr(qt, f, data[f])
    qt.total = (qt.amount or 0) + (qt.vat or 0) + (qt.delivery_cost or 0)
    if "items" in data:
        QuotationItem.query.filter_by(quotation_id=qid).delete(); db.session.flush()
        for it in data["items"]:
            qty = float(it.get("quantity",0)); up = float(it.get("unit_price",0))
            db.session.add(QuotationItem(quotation_id=qid, item_name=it.get("item_name",""),
                quantity=qty, unit_price=up, total=round(qty*up,2)))
    AuditService.log("edit","quotation",qid,f"تعديل عرض سعر")
    db.session.commit()
    return ok(qt.to_dict())

@api.route("/procurement/quotations/<int:qid>/accept", methods=["POST"])
@require_role("admin","manager")
def accept_quotation(qid):
    qt = Quotation.query.get_or_404(qid)
    qt.status = "accepted"
    AuditService.log("accept","quotation",qid,f"قبول عرض سعر")
    db.session.commit()
    return ok(qt.to_dict())

@api.route("/procurement/quotations/<int:qid>", methods=["DELETE"])
@require_role("admin")
def delete_quotation(qid):
    qt = Quotation.query.get_or_404(qid)
    db.session.delete(qt)
    AuditService.log("delete","quotation",qid,f"حذف عرض سعر")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — QUOTATION COMPARISON
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/rfq/<int:rid>/compare", methods=["GET"])
@jwt_required()
def compare_quotations(rid):
    rfq = RFQ.query.get_or_404(rid)
    qts = Quotation.query.filter_by(rfq_id=rid).all()
    items_map = {}
    for qt in qts:
        for it in qt.items:
            k = it.item_name
            if k not in items_map: items_map[k] = []
            items_map[k].append({"quotation_id":qt.id,"supplier_name":qt.supplier.name if qt.supplier else "","item_name":k,
                "quantity":it.quantity,"unit_price":it.unit_price,"total":it.total})
    best = None
    for qt in qts:
        score = (5 if qt.status=="accepted" else 0) + max(0, 5 - (qt.delivery_days or 99)/10)
        if not best or score > best["score"]:
            best = {"id":qt.id,"supplier_name":qt.supplier.name if qt.supplier else "","total":qt.total,"delivery_days":qt.delivery_days,"score":score}
    return ok({"rfq":rfq.to_dict(),"quotations":[q.to_dict() for q in qts],"comparison":items_map,"best_offer":best})

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — PURCHASE ORDERS
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/po", methods=["GET"])
@jwt_required()
def get_po_list():
    search = request.args.get("search",""); status = request.args.get("status","")
    q = PurchaseOrder.query
    if search: q = q.filter(PurchaseOrder.ref_number.contains(search))
    if status: q = q.filter(PurchaseOrder.status == status)
    return ok([p.to_dict() for p in q.order_by(PurchaseOrder.created_at.desc()).all()])

@api.route("/procurement/po", methods=["POST"])
@require_role("admin","manager")
def create_po():
    data = request.get_json() or {}
    ref = gen_ref("PO")
    po = PurchaseOrder(ref_number=ref, supplier_id=data.get("supplier_id"),
        warehouse_id=data.get("warehouse_id"), currency=data.get("currency","SAR"),
        payment_terms=data.get("payment_terms",""), notes=data.get("notes",""),
        created_by=g.current_user.id, status="draft")
    if "quotation_id" in data: po.quotation_id = int(data["quotation_id"])
    if "project_id" in data: po.project_id = int(data["project_id"])
    if "order_date" in data: po.order_date = parse_date(data["order_date"])
    if "delivery_date" in data: po.delivery_date = parse_date(data["delivery_date"])
    db.session.add(po); db.session.flush()
    for it in (data.get("items") or []):
        qty = float(it.get("quantity",0)); up = float(it.get("unit_price",0))
        db.session.add(POItem(po_id=po.id, item_name=it.get("item_name",""),
            quantity=qty, unit_price=up, total=round(qty*up,2)))
    AuditService.log("create","purchase_order",po.id,f"إنشاء أمر شراء: {ref}")
    db.session.commit()
    return created(po.to_dict())

@api.route("/procurement/po/<int:pid>", methods=["GET"])
@jwt_required()
def get_po(pid):
    po = PurchaseOrder.query.get_or_404(pid)
    return ok(po.to_dict())

@api.route("/procurement/po/<int:pid>", methods=["PUT"])
@require_role("admin","manager")
def update_po(pid):
    po = PurchaseOrder.query.get_or_404(pid)
    data = request.get_json() or {}
    for f in ("currency","payment_terms","notes","status"):
        if f in data: setattr(po, f, data[f])
    if "warehouse_id" in data: po.warehouse_id = int(data["warehouse_id"])
    if "delivery_date" in data: po.delivery_date = parse_date(data["delivery_date"])
    if "items" in data:
        POItem.query.filter_by(po_id=pid).delete(); db.session.flush()
        for it in data["items"]:
            qty = float(it.get("quantity",0)); up = float(it.get("unit_price",0))
            db.session.add(POItem(po_id=pid, item_name=it.get("item_name",""),
                quantity=qty, unit_price=up, total=round(qty*up,2)))
    AuditService.log("edit","purchase_order",pid,f"تعديل أمر شراء: {po.ref_number}")
    db.session.commit()
    return ok(po.to_dict())

@api.route("/procurement/po/<int:pid>/approve", methods=["POST"])
@require_role("admin","manager")
def approve_po(pid):
    po = PurchaseOrder.query.get_or_404(pid)
    if po.status != "draft": return err("يمكن اعتماد المسودات فقط")
    is_ok, msg, final = _process_chain_approval(po, int(get_jwt_identity()), "po")
    if not is_ok: return err(msg)
    if final in ("step", "step_same"):
        db.session.commit()
        return ok(po.to_dict(), msg)
    if final == "final":
        po.status = "approved"
        AuditService.log("approve","purchase_order",pid,f"اعتماد أمر شراء: {po.ref_number}")
        db.session.commit()
        return ok(po.to_dict())
    po.status = "approved"
    level = (request.get_json() or {}).get("level", 1)
    Appr = ApprovalLog(resource_type="po", resource_id=po.id,
           approver_id=g.current_user.id, status="approved", level=level)
    db.session.add(Appr)
    AuditService.log("approve","purchase_order",pid,f"اعتماد أمر شراء: {po.ref_number}")
    db.session.commit()
    return ok(po.to_dict())

@api.route("/procurement/po/<int:pid>/reject", methods=["POST"])
@require_role("admin","manager")
def reject_po(pid):
    po = PurchaseOrder.query.get_or_404(pid)
    if po.status != "draft": return err("يمكن رفض المسودات فقط")
    po.status = "cancelled"
    data = request.get_json() or {}
    Appr = ApprovalLog(resource_type="po", resource_id=po.id,
           approver_id=g.current_user.id, status="rejected", notes=data.get("notes",""))
    db.session.add(Appr)
    AuditService.log("reject","purchase_order",pid,f"رفض أمر شراء: {po.ref_number}")
    db.session.commit()
    return ok(po.to_dict())

@api.route("/procurement/po/<int:pid>/send", methods=["POST"])
@require_role("admin","manager")
def send_po(pid):
    po = PurchaseOrder.query.get_or_404(pid)
    if po.status not in ("draft","approved"): return err("يمكن إرسال المسودات أو المعتمدة فقط")
    po.status = "sent"
    AuditService.log("send","purchase_order",pid,f"إرسال أمر شراء: {po.ref_number}")
    db.session.commit()
    return ok(po.to_dict())

@api.route("/procurement/po/<int:pid>/cancel", methods=["POST"])
@require_role("admin","manager")
def cancel_po(pid):
    po = PurchaseOrder.query.get_or_404(pid)
    if po.status in ("completed","cancelled"): return err("لا يمكن إلغاء أمر شراء مكتمل أو ملغي")
    po.status = "cancelled"
    AuditService.log("cancel","purchase_order",pid,f"إلغاء أمر شراء: {po.ref_number}")
    db.session.commit()
    return ok(po.to_dict())

@api.route("/procurement/po/<int:pid>", methods=["DELETE"])
@require_role("admin")
def delete_po(pid):
    po = PurchaseOrder.query.get_or_404(pid)
    if po.status in ("sent","partial","completed"): return err("لا يمكن حذف أمر شراء مرسل أو مكتمل")
    db.session.delete(po)
    AuditService.log("delete","purchase_order",pid,f"حذف أمر شراء")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — GOODS RECEIPT (GRN)
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/grn", methods=["GET"])
@jwt_required()
def get_grn_list():
    search = request.args.get("search","")
    q = GoodsReceipt.query
    if search: q = q.filter(GoodsReceipt.ref_number.contains(search))
    return ok([g.to_dict() for g in q.order_by(GoodsReceipt.created_at.desc()).all()])

@api.route("/procurement/grn", methods=["POST"])
@require_role("admin","manager","keeper")
def create_grn():
    data = request.get_json() or {}
    ref = gen_ref("GRN")
    grn = GoodsReceipt(ref_number=ref, po_id=data.get("po_id"),
        warehouse_id=data.get("warehouse_id"), supplier_id=data.get("supplier_id"),
        received_by=g.current_user.id, notes=data.get("notes",""))
    db.session.add(grn); db.session.flush()
    total_value = 0
    for it in (data.get("items") or []):
        ordered = float(it.get("ordered_qty",0))
        damaged = float(it.get("damaged_qty",0))
        rejected = float(it.get("rejected_qty",0))
        accepted = float(it.get("accepted_qty",ordered - damaged - rejected))
        unit_price = float(it.get("unit_price",0))
        total = round(accepted * unit_price, 2)
        total_value += total
        item_id = it.get("item_id")
        if not item_id and it.get("item_name"):
            item = Item.query.filter_by(name=it["item_name"]).first()
            if item: item_id = item.id
        gi = GRNItem(grn_id=grn.id, po_item_id=it.get("po_item_id"),
            item_id=item_id, item_name=it.get("item_name",""), ordered_qty=ordered,
            received_qty=float(it.get("received_qty",ordered)),
            damaged_qty=damaged, rejected_qty=rejected,
            accepted_qty=accepted, unit_price=unit_price, total=total)
        db.session.add(gi)
        # create StockMovement record for audit trail
        if accepted > 0 and grn.warehouse_id and item_id:
            try:
                MovementService.create({"item_id":item_id,"warehouse_id":grn.warehouse_id,
                    "quantity":accepted,"type":"in","reference":ref,"unit_price":unit_price,
                    "notes":f"GRN: {ref}"}, g.current_user.id)
            except Exception:
                pass  # non-blocking
        # update stock for accepted items
        if accepted > 0 and grn.warehouse_id and item_id:
            stk = Stock.query.filter_by(item_id=item_id, warehouse_id=grn.warehouse_id).first()
            if stk: stk.quantity += accepted
            else: db.session.add(Stock(item_id=item_id, warehouse_id=grn.warehouse_id, quantity=accepted))
            InventoryLayer.add_layer(item_id, grn.warehouse_id, accepted, unit_price, ref_type="grn", ref_id=grn.id)
    # update PO received quantities
    if grn.po_id:
        po = PurchaseOrder.query.get(grn.po_id)
        if po:
            grn_items = data.get("items") or []
            for poi in po.items:
                for gi_data in grn_items:
                    if gi_data.get("po_item_id") == poi.id or gi_data.get("item_name","") == poi.item_name:
                        poi.received_qty = (poi.received_qty or 0) + float(gi_data.get("accepted_qty",0))
            # update PO status
            all_complete = all((poi.received_qty or 0) >= (poi.quantity or 0) for poi in po.items)
            any_received = any((poi.received_qty or 0) > 0 for poi in po.items)
            if all_complete: po.status = "completed"
            elif any_received: po.status = "partial"
    AuditService.log("create","goods_receipt",grn.id,f"إنشاء إذن استلام: {ref}")
    db.session.commit()
    sup = Supplier.query.get(data.get("supplier_id"))
    send_notif_email(f"📥 إذن استلام جديد: {ref}",
        f"تم إستلام بضاعة:\nالمرجع: {ref}\nالمورد: {sup.name if sup else '—'}\n"
        f"إجمالي القيمة: {total_value:,.2f} ر.س\nالمستلم: {_uname(g.current_user.id)}")
    return created(grn.to_dict())

@api.route("/procurement/grn/<int:gid>", methods=["GET"])
@jwt_required()
def get_grn(gid):
    grn = GoodsReceipt.query.get_or_404(gid)
    return ok(grn.to_dict())

@api.route("/procurement/grn/<int:gid>/cancel", methods=["POST"])
@require_role("admin")
def cancel_grn(gid):
    grn = GoodsReceipt.query.get_or_404(gid)
    for it in grn.items:
        if it.item_id and grn.warehouse_id:
            stk = Stock.query.filter_by(item_id=it.item_id, warehouse_id=grn.warehouse_id).first()
            if stk: stk.quantity = max(0, stk.quantity - (it.accepted_qty or 0))
            # reverse FIFO layers
            InventoryLayer.consume(it.item_id, grn.warehouse_id, it.accepted_qty or 0)
    if grn.po:
        for poi in grn.po.items:
            poi.received_qty = max(0, (poi.received_qty or 0) - sum(gi.quantity for gi in grn.items if gi.po_item_id == poi.id))
    AuditService.log("cancel","goods_receipt",gid,f"إلغاء إذن استلام: {grn.ref_number}")
    db.session.delete(grn)
    db.session.commit()
    return ok(message="تم إلغاء إذن الاستلام")

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — PURCHASE RETURNS
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/returns", methods=["GET"])
@jwt_required()
def get_preturn_list():
    search = request.args.get("search","")
    q = PurchaseReturn.query
    if search: q = q.filter(PurchaseReturn.ref_number.contains(search))
    return ok([r.to_dict() for r in q.order_by(PurchaseReturn.created_at.desc()).all()])

@api.route("/procurement/returns", methods=["POST"])
@require_role("admin","manager","keeper")
def create_preturn():
    data = request.get_json() or {}
    ref = gen_ref("SR")
    pr = PurchaseReturn(ref_number=ref, supplier_id=data.get("supplier_id"),
        warehouse_id=data.get("warehouse_id"), reason=data.get("reason","other"),
        notes=data.get("notes",""), status="pending")
    db.session.add(pr); db.session.flush()
    for it in (data.get("items") or []):
        qty = float(it.get("quantity",0)); up = float(it.get("unit_price",0))
        db.session.add(PReturnItem(return_id=pr.id, item_id=it.get("item_id"),
            item_name=it.get("item_name",""), quantity=qty, unit_price=up,
            total=round(qty*up,2), reason_detail=it.get("reason_detail","")))
    AuditService.log("create","purchase_return",pr.id,f"إنشاء مرتجع: {ref}")
    db.session.commit()
    return created(pr.to_dict())

@api.route("/procurement/returns/<int:rid>", methods=["GET"])
@jwt_required()
def get_preturn(rid):
    pr = PurchaseReturn.query.get_or_404(rid)
    return ok(pr.to_dict())

@api.route("/procurement/returns/<int:rid>/approve", methods=["POST"])
@require_role("admin","manager")
def approve_preturn(rid):
    pr = PurchaseReturn.query.get_or_404(rid)
    if pr.status != "pending": return err("يمكن اعتماد المعلقات فقط")
    pr.status = "returned"
    for it in pr.items:
        if it.item_id and pr.warehouse_id:
            stk = Stock.query.filter_by(item_id=it.item_id, warehouse_id=pr.warehouse_id).first()
            if stk: stk.quantity = max(0, stk.quantity - it.quantity)
    AuditService.log("approve","purchase_return",rid,f"اعتماد مرتجع: {pr.ref_number}")
    db.session.commit()
    return ok(pr.to_dict())

@api.route("/procurement/returns/<int:rid>/reject", methods=["POST"])
@require_role("admin","manager")
def reject_preturn(rid):
    pr = PurchaseReturn.query.get_or_404(rid)
    if pr.status != "pending": return err("يمكن رفض المعلقات فقط")
    pr.status = "cancelled"
    AuditService.log("reject","purchase_return",rid,f"رفض مرتجع: {pr.ref_number}")
    db.session.commit()
    return ok(pr.to_dict())

@api.route("/procurement/returns/<int:rid>", methods=["DELETE"])
@require_role("admin")
def delete_preturn(rid):
    pr = PurchaseReturn.query.get_or_404(rid)
    db.session.delete(pr)
    AuditService.log("delete","purchase_return",rid,f"حذف مرتجع")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — SUPPLIER EVALUATION
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/evaluations", methods=["GET"])
@jwt_required()
def get_evaluations():
    supplier_id = request.args.get("supplier_id")
    q = SupplierEvaluation.query
    if supplier_id: q = q.filter(SupplierEvaluation.supplier_id == int(supplier_id))
    return ok([e.to_dict() for e in q.order_by(SupplierEvaluation.created_at.desc()).all()])

@api.route("/procurement/evaluations", methods=["POST"])
@require_role("admin","manager")
def create_evaluation():
    data = request.get_json() or {}
    ev = SupplierEvaluation(supplier_id=data.get("supplier_id"), po_id=data.get("po_id"),
        delivery_accuracy=int(data.get("delivery_accuracy",3)),
        quality_score=int(data.get("quality_score",3)),
        response_speed=int(data.get("response_speed",3)),
        price_competitiveness=int(data.get("price_competitiveness",3)),
        notes=data.get("notes",""), evaluator_id=g.current_user.id)
    db.session.add(ev); db.session.flush()
    # update supplier profile averages
    sid = data.get("supplier_id")
    if sid:
        profile = SupplierProfile.query.filter_by(supplier_id=int(sid)).first()
        if not profile:
            profile = SupplierProfile(supplier_id=int(sid))
            db.session.add(profile)
        evals = SupplierEvaluation.query.filter_by(supplier_id=int(sid)).all()
        n = len(evals)
        profile.delivery_accuracy = round(sum(e.delivery_accuracy for e in evals) / n, 1)
        profile.quality_score = round(sum(e.quality_score for e in evals) / n, 1)
        profile.response_speed = round(sum(e.response_speed for e in evals) / n, 1)
        profile.price_competitiveness = round(sum(e.price_competitiveness for e in evals) / n, 1)
        profile.total_evaluations = n
    AuditService.log("create","supplier_evaluation",ev.id,f"تقييم مورد")
    db.session.commit()
    return created(ev.to_dict())

@api.route("/procurement/evaluations/<int:eid>", methods=["DELETE"])
@require_role("admin")
def delete_evaluation(eid):
    ev = SupplierEvaluation.query.get_or_404(eid)
    db.session.delete(ev)
    AuditService.log("delete","supplier_evaluation",eid,f"حذف تقييم مورد")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — SUPPLIER PROFILES
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/suppliers/<int:sid>/profile", methods=["GET"])
@jwt_required()
def get_supplier_profile(sid):
    profile = SupplierProfile.query.filter_by(supplier_id=sid).first()
    if not profile:
        profile = SupplierProfile(supplier_id=sid)
        db.session.add(profile); db.session.commit()
    return ok(profile.to_dict())

@api.route("/procurement/suppliers/<int:sid>/profile", methods=["PUT"])
@require_role("admin","manager")
def update_supplier_profile(sid):
    profile = SupplierProfile.query.filter_by(supplier_id=sid).first()
    if not profile:
        profile = SupplierProfile(supplier_id=sid)
        db.session.add(profile)
    data = request.get_json() or {}
    for f in ("contact_person","commercial_register","website"):
        if f in data: setattr(profile, f, data[f])
    AuditService.log("edit","supplier_profile",sid,f"تحديث بيانات مورد")
    db.session.commit()
    return ok(profile.to_dict())

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — DASHBOARD
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/dashboard", methods=["GET"])
@jwt_required()
def procurement_dashboard():
    now = datetime.datetime.utcnow()
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    pr_total = PurchaseRequest.query.count()
    pr_pending = PurchaseRequest.query.filter_by(status="pending").count()
    rfq_active = RFQ.query.filter(RFQ.status.in_(["draft","sent"])).count()
    po_total = PurchaseOrder.query.count()
    po_open = PurchaseOrder.query.filter(PurchaseOrder.status.in_(["draft","approved","sent","partial"])).count()
    po_completed = PurchaseOrder.query.filter_by(status="completed").count()
    # monthly purchases
    monthly_pos = PurchaseOrder.query.filter(PurchaseOrder.created_at >= month_start).all()
    monthly_total = sum(p.total_amount or 0 for p in monthly_pos)
    # supplier ranking
    top_suppliers = db.session.query(
        SupplierProfile.supplier_id, Supplier.name,
        (SupplierProfile.delivery_accuracy + SupplierProfile.quality_score +
         SupplierProfile.response_speed + SupplierProfile.price_competitiveness) / 4
    ).join(Supplier, Supplier.id == SupplierProfile.supplier_id
    ).order_by(SupplierProfile.total_evaluations.desc()).limit(10).all()
    return ok({
        "pr_total": pr_total, "pr_pending": pr_pending,
        "rfq_active": rfq_active,
        "po_total": po_total, "po_open": po_open, "po_completed": po_completed,
        "monthly_purchases": round(monthly_total, 2),
        "top_suppliers": [{"supplier_id":s[0],"name":s[1],"avg_score":round(float(s[2]),1)} for s in top_suppliers],
    })

# ══════════════════════════════════════════════════════════════
#  PROCUREMENT — REPORTS
# ══════════════════════════════════════════════════════════════
@api.route("/procurement/reports/pr", methods=["GET"])
@jwt_required()
def pr_report():
    from_date = request.args.get("from"); to_date = request.args.get("to")
    q = PurchaseRequest.query
    if from_date: q = q.filter(PurchaseRequest.created_at >= parse_date(from_date))
    if to_date: q = q.filter(PurchaseRequest.created_at <= parse_date(to_date) + datetime.timedelta(days=1))
    items = [p.to_dict() for p in q.order_by(PurchaseRequest.created_at.desc()).all()]
    total_est = sum(sum(it.get("estimated_cost",0)*it.get("quantity",0) for it in p.get("items",[])) for p in items)
    return ok({"items":items,"total":len(items),"total_estimated":round(total_est,2)})

@api.route("/procurement/reports/po", methods=["GET"])
@jwt_required()
def po_report():
    from_date = request.args.get("from"); to_date = request.args.get("to")
    q = PurchaseOrder.query
    if from_date: q = q.filter(PurchaseOrder.created_at >= parse_date(from_date))
    if to_date: q = q.filter(PurchaseOrder.created_at <= parse_date(to_date) + datetime.timedelta(days=1))
    items = [p.to_dict() for p in q.order_by(PurchaseOrder.created_at.desc()).all()]
    total = sum(p.total_amount or 0 for p in q)
    return ok({"items":items,"total":len(items),"total_amount":round(total,2)})

@api.route("/procurement/reports/suppliers", methods=["GET"])
@jwt_required()
def supplier_report():
    profiles = SupplierProfile.query.order_by(SupplierProfile.total_evaluations.desc()).all()
    return ok([p.to_dict() for p in profiles])

# ══════════════════════════════════════════════════════════════
#  SALES — Sale Orders
# ══════════════════════════════════════════════════════════════
@api.route("/sales", methods=["GET"])
@jwt_required()
def get_sales():
    search = request.args.get("search","")
    status = request.args.get("status","")
    q = SaleOrder.query
    if search: q = q.filter(SaleOrder.ref_number.contains(search) | SaleOrder.customer_name.contains(search))
    if status: q = q.filter(SaleOrder.status == status)
    return ok([s.to_dict() for s in q.order_by(SaleOrder.created_at.desc()).all()])

@api.route("/sales", methods=["POST"])
@require_role("admin","manager","keeper")
def create_sale():
    data = request.get_json() or {}
    if not data.get("customer_name"): return err("اسم العميل مطلوب")
    if not data.get("items"): return err("يجب إضافة أصناف على الأقل")
    ref = gen_ref("SO")
    sale = SaleOrder(ref_number=ref, customer_name=data["customer_name"],
        customer_phone=data.get("customer_phone",""),
        customer_email=data.get("customer_email",""),
        warehouse_id=data.get("warehouse_id"),
        discount_pct=float(data.get("discount_pct",0)),
        tax_pct=float(data.get("tax_pct",0)),
        notes=data.get("notes",""), status="pending",
        created_by=g.current_user.id)
    if "sale_date" in data: sale.sale_date = parse_date(data["sale_date"])
    db.session.add(sale); db.session.flush()
    subtotal = 0
    for it in (data.get("items") or []):
        qty = float(it.get("quantity",0)); up = float(it.get("unit_price",0))
        total = round(qty * up, 2)
        item_id = it.get("item_id")
        if not item_id:
            item = Item.query.filter_by(name=it.get("item_name","")).first()
            if item: item_id = item.id
        db.session.add(SaleItem(sale_id=sale.id, item_id=item_id,
            item_name=it.get("item_name",""), quantity=qty,
            unit_price=up, total=total))
        subtotal += total
    discount_amt = round(subtotal * sale.discount_pct / 100, 2) if sale.discount_pct else 0
    tax_amt = round((subtotal - discount_amt) * sale.tax_pct / 100, 2) if sale.tax_pct else 0
    sale.subtotal = subtotal; sale.discount_amt = discount_amt
    sale.tax_amt = tax_amt; sale.total = round(subtotal - discount_amt + tax_amt, 2)
    AuditService.log("create","sale_order",sale.id,f"إنشاء أمر بيع: {ref}")
    db.session.commit()
    return created(sale.to_dict())

@api.route("/sales/<int:sid>", methods=["GET"])
@jwt_required()
def get_sale(sid):
    sale = SaleOrder.query.get_or_404(sid)
    return ok(sale.to_dict())

@api.route("/sales/<int:sid>", methods=["PUT"])
@require_role("admin","manager")
def update_sale(sid):
    sale = SaleOrder.query.get_or_404(sid)
    if sale.status not in ("pending",): return err("يمكن تعديل الطلبات المعلقة فقط")
    data = request.get_json() or {}
    for f in ("customer_name","customer_phone","customer_email","notes"):
        if f in data: setattr(sale, f, data[f])
    if "discount_pct" in data: sale.discount_pct = float(data["discount_pct"])
    if "tax_pct" in data: sale.tax_pct = float(data["tax_pct"])
    if "warehouse_id" in data: sale.warehouse_id = int(data["warehouse_id"])
    if "items" in data:
        SaleItem.query.filter_by(sale_id=sid).delete(); db.session.flush()
        subtotal = 0
        for it in data["items"]:
            qty = float(it.get("quantity",0)); up = float(it.get("unit_price",0))
            tot = round(qty * up, 2)
            item_id = it.get("item_id")
            if not item_id:
                item = Item.query.filter_by(name=it.get("item_name","")).first()
                if item: item_id = item.id
            db.session.add(SaleItem(sale_id=sid, item_id=item_id,
                item_name=it.get("item_name",""), quantity=qty,
                unit_price=up, total=tot))
            subtotal += tot
        discount_amt = round(subtotal * sale.discount_pct / 100, 2) if sale.discount_pct else 0
        tax_amt = round((subtotal - discount_amt) * sale.tax_pct / 100, 2) if sale.tax_pct else 0
        sale.subtotal = subtotal; sale.discount_amt = discount_amt
        sale.tax_amt = tax_amt; sale.total = round(subtotal - discount_amt + tax_amt, 2)
    AuditService.log("edit","sale_order",sid,f"تعديل أمر بيع: {sale.ref_number}")
    db.session.commit()
    return ok(sale.to_dict())

@api.route("/sales/<int:sid>/confirm", methods=["POST"])
@require_role("admin","manager")
def confirm_sale(sid):
    sale = SaleOrder.query.get_or_404(sid)
    if sale.status != "pending": return err("يمكن تأكيد الطلبات المعلقة فقط")
    if not sale.warehouse_id: return err("يجب تحديد المخزن")
    # deduct stock for each item
    for si in sale.items:
        if not si.item_id: continue
        stk = Stock.query.filter_by(item_id=si.item_id, warehouse_id=sale.warehouse_id).first()
        if not stk or stk.quantity < si.quantity:
            return err(f"الكمية غير كافية للصنف: {si.item_name} (المتاح: {stk.quantity if stk else 0})")
    for si in sale.items:
        if not si.item_id: continue
        stk = Stock.query.filter_by(item_id=si.item_id, warehouse_id=sale.warehouse_id).first()
        if stk: stk.quantity = max(0, stk.quantity - si.quantity)
        m = StockMovement(ref_number=gen_ref("SO"), type="out", item_id=si.item_id,
            warehouse_id=sale.warehouse_id, quantity=si.quantity,
            unit_price=si.unit_price, user_id=g.current_user.id,
            notes=f"أمر بيع: {sale.ref_number}")
        db.session.add(m)
    sale.status = "confirmed"
    AuditService.log("confirm","sale_order",sid,f"تأكيد أمر بيع: {sale.ref_number}")
    db.session.commit()
    return ok(sale.to_dict())

@api.route("/sales/<int:sid>/invoice", methods=["GET"])
@jwt_required()
def sale_invoice(sid):
    sale = SaleOrder.query.get_or_404(sid)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
    except ImportError:
        return err("reportlab غير مثبت")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=1.5*cm, leftMargin=1.5*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    ts = ParagraphStyle("T", parent=styles["Title"], fontSize=16, alignment=TA_CENTER)
    company = "شركة النبلاء للمقاولات العامة والإنشاءات"
    elems = [Paragraph(f"<b>{company}</b>", ts),
             Paragraph(f"فاتورة بيع — {sale.ref_number}", styles["Normal"]),
             Spacer(1, 0.3*cm)]
    info = [
        ["رقم الفاتورة:", sale.ref_number],
        ["التاريخ:", sale.sale_date.isoformat() if sale.sale_date else ""],
        ["العميل:", sale.customer_name],
        ["الهاتف:", sale.customer_phone or "—"],
        ["المخزن:", sale.warehouse.name if sale.warehouse else "—"],
        ["الحالة:", sale.STATUS_LABELS.get(sale.status, sale.status)],
    ]
    t = Table(info, colWidths=[4*cm, 10*cm])
    t.setStyle(TableStyle([
        ("FONTSIZE",(0,0),(-1,-1),10),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#BDC3C7")),
        ("BACKGROUND",(0,0),(0,-1),colors.HexColor("#1B4F72")),
        ("TEXTCOLOR",(0,0),(0,-1),colors.white),
    ]))
    elems.append(t); elems.append(Spacer(1,0.3*cm))
    hdr = ["الصنف","الكمية","سعر الوحدة","الإجمالي"]
    rows = [[si.item_name, str(si.quantity), f"{si.unit_price:,.2f}", f"{si.total:,.2f}"] for si in sale.items]
    rows.append(["","","الإجمالي:", f"{sale.subtotal:,.2f}"])
    if sale.discount_amt:
        rows.append(["","","الخصم:", f"-{sale.discount_amt:,.2f}"])
    if sale.tax_amt:
        rows.append(["","","الضريبة:", f"{sale.tax_amt:,.2f}"])
    rows.append(["","","<b>المجموع النهائي:</b>", f"<b>{sale.total:,.2f}</b>"])
    tt = Table([hdr]+rows, colWidths=[6*cm,3*cm,4*cm,4*cm], repeatRows=1)
    tt.setStyle(TableStyle([
        ("FONTSIZE",(0,0),(-1,-1),9),("ALIGN",(1,0),(-1,-1),"CENTER"),
        ("GRID",(0,0),(-1,-2),0.3,colors.HexColor("#BDC3C7")),
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1B4F72")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("LINEBELOW",(0,-1),(-1,-1),1,colors.HexColor("#1B4F72")),
    ]))
    elems.append(tt)
    doc.build(elems); buf.seek(0)
    return send_file(buf, mimetype="application/pdf",
                     download_name=f"sale_{sale.ref_number}.pdf", as_attachment=True)

@api.route("/sales/<int:sid>/cancel", methods=["POST"])
@require_role("admin","manager")
def cancel_sale(sid):
    sale = SaleOrder.query.get_or_404(sid)
    if sale.status in ("cancelled",): return err("الطلب ملغي بالفعل")
    # reverse stock if confirmed
    if sale.status == "confirmed" and sale.warehouse_id:
        for si in sale.items:
            if not si.item_id: continue
            stk = Stock.query.filter_by(item_id=si.item_id, warehouse_id=sale.warehouse_id).first()
            if stk: stk.quantity += si.quantity
    sale.status = "cancelled"
    AuditService.log("cancel","sale_order",sid,f"إلغاء أمر بيع: {sale.ref_number}")
    db.session.commit()
    return ok(sale.to_dict())

@api.route("/sales/<int:sid>", methods=["DELETE"])
@require_role("admin")
def delete_sale(sid):
    sale = SaleOrder.query.get_or_404(sid)
    if sale.status == "confirmed": return err("لا يمكن حذف أمر مؤكد")
    db.session.delete(sale)
    AuditService.log("delete","sale_order",sid,f"حذف أمر بيع")
    db.session.commit()
    return ok(message="تم الحذف")

# ══════════════════════════════════════════════════════════════
#  SUPPLIER PORTAL
# ══════════════════════════════════════════════════════════════
@api.route("/supplier/login", methods=["POST"])
def supplier_login():
    data = request.get_json() or {}
    username = data.get("username", "").strip()
    password = data.get("password", "")
    if not username or not password: return err("اسم المستخدم وكلمة المرور مطلوبان")
    user = User.query.filter_by(username=username).first()
    if not user or not check_password_hash(user.password_hash, password): return err("بيانات الدخول غير صحيحة")
    if not user.is_active: return err("الحساب غير نشط")
    if not user.supplier_id: return err("هذا الحساب ليس لمورد")
    supplier = Supplier.query.get(user.supplier_id)
    if not supplier or not supplier.is_active: return err("المورد غير نشط")
    access_token = create_access_token(identity=str(user.id), additional_claims={"supplier_id": user.supplier_id, "role": "supplier"})
    return ok({"access_token": access_token, "user": {"id": user.id, "name": user.name, "username": user.username, "supplier_id": user.supplier_id, "supplier_name": supplier.name}})

def _supplier_required():
    uid = int(get_jwt_identity())
    user = User.query.get(uid)
    if not user or not user.supplier_id: return None
    return user

@api.route("/supplier/profile", methods=["GET"])
@jwt_required()
def supplier_get_profile():
    user = _supplier_required()
    if not user: return err("غير مصرح", 403)
    s = Supplier.query.get(user.supplier_id)
    if not s: return err("المورد غير موجود")
    p = SupplierProfile.query.filter_by(id=user.supplier_id).first()
    return ok({"id": s.id, "name": s.name, "code": s.code, "phone": s.phone, "email": s.email, "address": s.address, "tax_number": s.tax_number, "payment_terms": s.payment_terms, "rating": p.total_evaluations if p else 0, "delivery_accuracy": p.delivery_accuracy if p else 0, "quality_score": p.quality_score if p else 0})

@api.route("/supplier/profile", methods=["PUT"])
@jwt_required()
def supplier_update_profile():
    user = _supplier_required()
    if not user: return err("غير مصرح", 403)
    s = Supplier.query.get(user.supplier_id)
    if not s: return err("المورد غير موجود")
    data = request.get_json() or {}
    for f in ["phone", "email", "address", "tax_number", "payment_terms"]:
        if f in data: setattr(s, f, data[f])
    db.session.commit()
    return ok(message="✅ تم تحديث البيانات")

@api.route("/supplier/rfqs", methods=["GET"])
@jwt_required()
def supplier_rfqs():
    user = _supplier_required()
    if not user: return err("غير مصرح", 403)
    rfqs = RFQ.query.filter(
        RFQ.status.in_(["sent", "open"]),
        RFQ.id.in_(db.session.query(RFQSupplier.rfq_id).filter(RFQSupplier.supplier_id == user.supplier_id))
    ).order_by(RFQ.created_at.desc()).all()
    result = []
    for r in rfqs:
        rs = RFQSupplier.query.filter_by(rfq_id=r.id, supplier_id=user.supplier_id).first()
        existing_q = Quotation.query.filter_by(rfq_id=r.id, supplier_id=user.supplier_id).first()
        result.append({"id": r.id, "ref_number": r.ref_number, "valid_until": r.valid_until.isoformat() if r.valid_until else "", "delivery_terms": r.delivery_terms, "notes": r.notes, "status": r.status, "created_at": r.created_at.isoformat(), "responded": rs.responded if rs else False, "has_quotation": existing_q is not None, "quotation_id": existing_q.id if existing_q else None, "items": [{"id": qi.id, "item_id": qi.item_id, "item_name": qi.item_name, "quantity": qi.quantity, "unit": qi.unit} for qi in r.items]})
    return ok(result)

@api.route("/supplier/rfqs/<int:rid>/quote", methods=["POST"])
@jwt_required()
def supplier_submit_quote(rid):
    user = _supplier_required()
    if not user: return err("غير مصرح", 403)
    rfq = RFQ.query.get_or_404(rid)
    existing = Quotation.query.filter_by(rfq_id=rid, supplier_id=user.supplier_id).first()
    if existing: return err("سبق تقديم عرض سعر لهذا الطلب")
    rs = RFQSupplier.query.filter_by(rfq_id=rid, supplier_id=user.supplier_id).first()
    if not rs: return err("أنت غير مدرج في هذا الطلب")
    data = request.get_json() or {}
    items = data.get("items", [])
    if not items: return err("يجب إضافة أصناف على الأقل")
    q = Quotation(rfq_id=rid, supplier_id=user.supplier_id, ref_number=gen_ref("Q"), amount=data.get("amount", 0), delivery_days=data.get("delivery_days"), warranty_period=data.get("warranty_period"), valid_until=data.get("valid_until"), notes=data.get("notes"), status="submitted")
    db.session.add(q); db.session.flush()
    for it in items:
        db.session.add(QuotationItem(quotation_id=q.id, item_name=it.get("item_name", ""), quantity=it.get("quantity", 0), unit_price=it.get("unit_price", 0), total=it.get("quantity", 0) * it.get("unit_price", 0)))
    rs.responded = True
    db.session.commit()
    return created(q.to_dict(), message="✅ تم تقديم عرض السعر")

@api.route("/supplier/orders", methods=["GET"])
@jwt_required()
def supplier_orders():
    user = _supplier_required()
    if not user: return err("غير مصرح", 403)
    orders = PurchaseOrder.query.filter_by(supplier_id=user.supplier_id).order_by(PurchaseOrder.created_at.desc()).all()
    return ok([o.to_dict() for o in orders])

@api.route("/supplier/evaluations", methods=["GET"])
@jwt_required()
def supplier_evaluations():
    user = _supplier_required()
    if not user: return err("غير مصرح", 403)
    evals = SupplierEvaluation.query.filter_by(supplier_id=user.supplier_id).order_by(SupplierEvaluation.created_at.desc()).all()
    return ok([e.to_dict() for e in evals])

@api.route("/supplier/documents", methods=["POST"])
@jwt_required()
def supplier_upload_doc():
    user = _supplier_required()
    if not user: return err("غير مصرح", 403)
    if "file" not in request.files: return err("اختر ملفاً")
    f = request.files["file"]
    if f.filename == "": return err("اسم الملف فارغ")
    ext = f.filename.rsplit(".", 1)[-1].lower() if "." in f.filename else ""
    if ext not in current_app.config.get("ALLOWED_EXTENSIONS", {"pdf","jpg","jpeg","png","doc","docx"}): return err("نوع الملف غير مسموح")
    filename = secure_filename(f"supplier_{user.supplier_id}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}_{f.filename}")
    upload_dir = os.path.join(os.path.dirname(current_app.root_path), "uploads", "documents")
    os.makedirs(upload_dir, exist_ok=True)
    f.save(os.path.join(upload_dir, filename))
    AuditService.log("upload", "supplier_document", user.supplier_id, f"مستند: {f.filename}")
    return ok({"filename": filename}, message="✅ تم رفع المستند")

# ══════════════════════════════════════════════════════════════
#  AI ASSISTANT
# ══════════════════════════════════════════════════════════════
@api.route("/ai/query", methods=["POST"])
@jwt_required()
def ai_query():
    data = request.get_json() or {}
    q = (data.get("query") or "").strip()
    if not q: return err("السؤال مطلوب")
    ql = q.lower()
    # Specific intents first
    if any(w in ql for w in ["حركة","حركات","movement","صرف","إدخال","وارد","أخر"]):
        recent = StockMovement.query.order_by(StockMovement.created_at.desc()).limit(5).all()
        lines = "\n".join([f"• {'🟢' if m.type in ('in','transfer_in') else '🔴'} {m.type} {m.quantity:.0f} {m.item.name if m.item else ''} ({m.created_at.strftime('%m/%d')})" for m in recent])
        return ok({"answer":f"آخر 5 حركات:\n{lines}"})
    if any(w in ql for w in ["مبيعات","sales","بيع","فاتورة"]):
        total_sales = db.session.query(func.sum(SaleOrder.total)).filter(SaleOrder.status!="cancelled").scalar() or 0
        count = SaleOrder.query.filter(SaleOrder.status!="cancelled").count()
        return ok({"answer":f"إجمالي المبيعات: {count} فاتورة بقيمة {total_sales:,.2f} ر.س"})
    if any(w in ql for w in ["منخفض","critical","حرج","نفذ","low","ناقص"]):
        all_items = Item.query.filter(Item.min_quantity > 0).all()
        crit = [i for i in all_items if i.get_total_stock() < i.min_quantity]
        names = "\n".join([f"• {i.name} ({i.get_total_stock():.0f}/{i.min_quantity:.0f})" for i in crit[:10]])
        return ok({"answer":f"أصناف منخفضة ({len(crit)}):\n{names or '— لا يوجد'}"})
    if any(w in ql for w in ["صنف","item","items","أصناف"]):
        items = Item.query.filter_by(is_active=True).count()
        total_stock = db.session.query(func.sum(Stock.quantity)).scalar() or 0
        return ok({"answer":f"إجمالي الأصناف النشطة: {items} صنف | إجمالي المخزون: {total_stock:,.0f} وحدة"})
    # Generic totals
    items = Item.query.filter_by(is_active=True).count()
    total_stock = db.session.query(func.sum(Stock.quantity)).scalar() or 0
    return ok({"answer":f"إجمالي الأصناف: {items} | المخزون: {total_stock:,.0f} وحدة"})
    return ok({"answer":"أستطيع الإجابة عن: المخزون، الأصناف، الحركات، المبيعات، الأصناف المنخفضة. جرب سؤالاً!"})

@api.route("/ai/analyze", methods=["GET"])
@jwt_required()
def ai_analyze():
    items = Item.query.filter_by(is_active=True).count()
    total_stock = db.session.query(func.sum(Stock.quantity)).scalar() or 0
    stocks = Stock.query.options(db.joinedload(Stock.item)).all()
    total_val = sum(s.quantity * (s.item.unit_price or 0) for s in stocks if s.item)
    all_its = Item.query.filter(Item.min_quantity > 0).all()
    critical = sum(1 for i in all_its if i.get_total_stock() < i.min_quantity)
    sales = db.session.query(func.sum(SaleOrder.total)).filter(SaleOrder.status!="cancelled").scalar() or 0
    return ok({
        "summary":f"📊 {items} صنف، {total_stock:,.0f} وحدة، {total_val:,.2f} ر.س قيمة، {critical} حرج، {sales:,.2f} ر.س مبيعات",
        "items":items,"total_stock":total_stock,"total_value":total_val,
        "critical":critical,"sales_total":sales,
        "recommendations":[
            "📦 راجع الأصناف منخفضة المخزون لتجنب نفادها",
            "🔄 جدول جرد دوري للأصناف عالية القيمة",
            "💰 راجع المبيعات الشهرية لتحديد الأصناف الأكثر طلباً",
        ]})

# ══════════════════════════════════════════════════════════════
#  OCR — Document Scanning
# ══════════════════════════════════════════════════════════════
@api.route("/ocr/scan", methods=["POST"])
@jwt_required()
def ocr_scan():
    if "file" not in request.files: return err("اختر ملف صورة")
    f = request.files["file"]
    if f.filename == "": return err("اسم الملف فارغ")
    try:
        from PIL import Image
        import pytesseract
    except ImportError:
        return err("مكتبة OCR غير مثبتة. قم بتشغيل: pip install pytesseract Pillow")
    try:
        img = Image.open(f)
        text = pytesseract.image_to_string(img, lang="ara+eng")
        lines = [l.strip() for l in text.split("\n") if l.strip()]
        return ok({"text":text,"lines":lines,"word_count":len(text.split())})
    except Exception as e:
        return err(f"فشل OCR: {str(e)[:100]}")
